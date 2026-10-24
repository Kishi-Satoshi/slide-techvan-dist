"""計画・データの型（グラフ・スケジュール・体制・費用・リスク・進捗）。

chart / schedule / team / cost / risks / status を登録する。
renderers.py が自動で読み込む。ここでは配置だけを決め、色・角丸・影などの見た目は
style（styles_impl.py）の共通の部品と色の役割（style.c）に任せる。
色は白（FFFFFF）以外は直接書かず、役割の色から作る（混ぜて淡くするのはグラフの淡い系列だけ）。
型で共通の決まり（styles_impl.py の BaseStyle「型で共通に使う決まり」）に従う:
本体は style.body() の Area の上端（a.y）から上詰めで置き、余りは下に残す。
文字の大きさは style.size() の段階（item_head / body / body_large / body_small / caption）だけを使い、
表の仲間（スケジュール・費用の表）は表の型と同じ table_head / table_body にそろえる。
番号の丸は style.marker、横に並べる列は a.x〜a.right を style.gutter() の間隔で分ける。
一覧（スケジュール・リスク・進捗）は、本文の枠（callout・note を除いた高さ）から
行の高さを決め、入りきらないときは詰めてから警告する。
"""
import math
import re

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION,
                             XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK)
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import slidekit as K
from styles_impl import Area
from renderers import register, as_items, check_count, render_table

WHITE = "FFFFFF"


# ====================================================================== 小さな部品（この型だけで使う）

def _z(style):
    """文字の大きさの段階（テーマの metrics.sizes）。型の中では大きさの数字を直接書かない。"""
    return {k: style.size(k, d) for k, d in (("item_head", 16), ("body", 14), ("body_large", 16),
                                             ("body_small", 12), ("caption", 11))}


def _table_sizes(style, z):
    """表の見出し・本文の大きさ（表の型 render_table と同じ段階）。"""
    return style.size("table_head", z["body_small"]), style.size("table_body", z["body_small"])


def _line_h(size):
    """1行の高さ（インチ）。"""
    return size * K.LINE_HEIGHT / 72


def _pill_h(size):
    """ラベル（pill・状況のラベル）の高さ。内側の余白をとっても文字が縮まない高さ。"""
    return _line_h(size) + 0.14


SAFE = 0.95   # 文字幅の見積もりに持たせる余裕（実際のフォントは見積もりより少し広いことがある）


def _balanced_w(text, w, size, bold=False):
    """折り返すときに各行の長さがそろう文字の枠の幅（1〜2 文字だけ次の行に送らない）。"""
    plain = K.plain(K.as_text(text))
    n = K.count_lines(plain, w * SAFE, size, bold)
    if n <= 1:
        return w
    bw = (math.ceil(K.text_em(plain, bold) / n) + 1.5) * size / 72 / SAFE
    return w if bw >= w or K.count_lines(plain, bw * SAFE, size, bold) > n else bw


def _mix(a, b, t):
    """色 a に色 b を t（0〜1）の割合で混ぜる（グラフの淡い系列を役割の色から作る）。"""
    a, b = a.lstrip("#"), b.lstrip("#")
    return "".join(f"{round(int(a[i:i + 2], 16) * (1 - t) + int(b[i:i + 2], 16) * t):02X}"
                   for i in (0, 2, 4))


def _lum(color):
    """相対輝度（WCAG の定義。0〜1）。"""
    h = color.lstrip("#")
    r, g, b = ((v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4)
               for v in (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _contrast(a, b):
    """2色のコントラスト比（1〜21）。"""
    hi, lo = sorted((_lum(a), _lum(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def _dark(fill):
    """白い地の上でもはっきり見える濃い色か（白い文字を載せても読めるか）。"""
    return _contrast(fill, WHITE) >= 3.0


def _on(c, fill):
    """塗りの上に載せる文字の色。白と本文の色のうち読みやすいほう（同じくらいなら白）。"""
    return WHITE if _contrast(fill, WHITE) * 1.15 >= _contrast(fill, c["text"]) else c["text"]


def _ink(c, color, bg, fallback=None):
    """淡い地に置く文字の色。地との差が 4.5:1 に届かなければ fallback（既定は本文の色）にする。"""
    return color if _contrast(color, bg) >= 4.5 else (fallback or c["text"])


def _fill_of(shp, default):
    """図形の塗りの色（読めなければ default）。"""
    try:
        return str(shp.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return default


def _recolor(shp, fill, line=None):
    """style の部品で描いた図形（角丸・影はそのまま）の塗りと線の色を変える。"""
    shp.fill.solid()
    shp.fill.fore_color.rgb = K.rgb(fill)
    shp.line.color.rgb = K.rgb(line or fill)
    return shp


def _attention(style, s, a, color):
    """目立たせる行（リスクの「高」・進捗の「遅延」）。強調の面（plain_card の highlight）に、
    その意味の色（テーマの levels・status の色）の線を添える。戻り値は面の色。"""
    c = style.c
    card = style.plain_card(s, a, highlight=True)
    card.line.color.rgb = K.rgb(color)
    card.line.width = Pt(1.25)
    return _fill_of(card, c.get("tint", c["bg_alt"]))


def _strong(c):
    """濃い塗り（体制図の責任者など）と、その上の文字色。"""
    return c["primary"], c["on_primary"]


def _grid(c):
    """罫線・目盛線・進捗の棒の地の淡い色（区切り線の色）。"""
    return c["rule"]


def _faint(c):
    """白い地と 3:1 ほどの差のある淡い色（線・矢印など文字でないもの）。"""
    return c.get("text_faint", c["text_muted"])


def _rename(items, key):
    """文字列だけの項目（{text: ...}）を、その型の項目名にそろえる。"""
    for it in items:
        if key not in it and "text" in it:
            it[key] = it.pop("text")
    return items


def _texts(v):
    """文字列のリストにそろえる（1つだけの文字列や {text: ...} の項目も受け付ける）。"""
    if not v:
        return []
    if isinstance(v, str):
        return [v]
    return [K.as_text(it.get("text", "")) for it in as_items(v)]


def _num(v):
    """'1,200' や 1200 を数にする。数でなければ None。"""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v if v is not None else "").replace(",", "").replace("，", "").strip()
    try:
        return float(t)
    except ValueError:
        return None


def _amount(v):
    """金額の表示。数字だけなら3桁ごとにカンマを入れる（'1200' → '1,200'）。"""
    t = str(v if v is not None else "").strip()
    m = re.fullmatch(r"([+\-−]?)(\d{4,})(\.\d+)?", t)
    if m:
        return f"{m.group(1)}{int(m.group(2)):,}{m.group(3) or ''}"
    return t


def _decimals(values):
    """小数点以下の桁数（最大2桁。0.1 未満の値があるときは 0 と表示されないよう4桁まで）。
    グラフの数の書式・費用の合計に使う。"""
    d = 0
    for v in values:
        if isinstance(v, str):
            v = _num(v)
        if isinstance(v, float) and not v.is_integer():
            k = 2 if abs(v) >= 0.1 else 4
            d = max(d, len(f"{v:.{k}f}".rstrip("0").split(".")[1]))
    return d


def _fmt(dec):
    """数の書式。小数があれば1桁目は必ず出し、2桁目は要るときだけ（28.4 / 17.25 / 9.0）。"""
    return "#,##0" + ("." + "0" + "#" * (dec - 1) if dec else "")


def _chip(style, s, x, y, w, h, text, fill, color, size):
    """塗りつぶしの小さなラベル（影響度・状況など、色が意味を持つもの）。"""
    b = K.add_box(s.shapes, x, y, w, h, fill=fill, geom="pill")
    style.text(s, 0, 0, 0, 0, text, size, color, bold=True, align="c", anchor="m", shape=b,
               margin=0.03, fit_min=9, label="ラベル")
    return b


def _chip_w(text, size, lo=0.6, hi=3.0):
    return min(hi, max(lo, K.text_em(K.plain(text), True) * size / 72 + 0.36))


def _head_h(z):
    """一覧の列見出しの行の高さ（next-actions と同じ）。"""
    return max(0.3, _line_h(z["caption"]) + 0.08)


def _head(style, s, x, y, w, h, text, align="l"):
    """一覧の列見出し（注釈の大きさの太字、控えめな色）。"""
    style.text(s, x, y, w, h, text, _z(style)["caption"], style.c["text_muted"], bold=True,
               align=align, anchor="m", fit_min=9, label="列見出し")


def _blank(style, s, x, y, w, h, size, align="c"):
    """値のないセル（担当・期限などがない行）。空けずに「—」を控えめな色で置く。"""
    style.text(s, x, y, w, h, "—", size, style.c["text_muted"], align=align, anchor="m")


def _person(s, x, y, d, fill, glyph):
    """人のアイコン（丸の中に頭と肩）。"""
    K.add_box(s.shapes, x, y, d, d, fill=fill, geom="oval")
    k = 0.5523 * 170
    head = [("M", 500, 190), ("C", 500 + k, 190, 670, 360 - k, 670, 360),
            ("C", 670, 360 + k, 500 + k, 530, 500, 530),
            ("C", 500 - k, 530, 330, 360 + k, 330, 360),
            ("C", 330, 360 - k, 500 - k, 190, 500, 190), ("Z",)]
    body = [("M", 190, 900), ("C", 190, 700, 330, 590, 500, 590),
            ("C", 670, 590, 810, 700, 810, 900), ("Z",)]
    K.add_path(s.shapes, x + d * 0.15, y + d * 0.13, d * 0.7, d * 0.7, [head, body], fill=glyph)


def _points_panel(style, s, a, points, title):
    """右側の「読み取りのポイント」。番号の丸（style.marker）と短い文を、上から同じ間隔で並べる。

    間隔は枠の高さを項目の数で等分した高さ（どの行も下に同じだけ余白が残る）。
    文が短くて間隔が広すぎるときは詰め、面もその分だけ短くする（余りは下に残す）。"""
    c = style.c
    z = _z(style)
    pad = 0.24
    head_y = a.y + 0.2
    hh = max(0.36, _line_h(z["item_head"]) + 0.04)
    top = head_y + hh + 0.18
    n = len(points)
    avail = a.bottom - 0.24 - top
    sep = 0.3 if n > 1 else 0.0          # 項目の間の最小の間（区切り線を引く）

    def layout(key, size):
        d = style.MARKER[key][0]
        tx = a.x + pad + d + 0.14
        tw = a.right - pad - tx
        lh = _line_h(size)
        off = max(0.0, (d - lh) / 2)      # 1行目の中心を丸の中心にそろえる
        # 見積もりより1行多く折り返しても区切り線にかからないよう、少し狭い幅で測る
        hs = [off + K.text_height(K.plain(p), (tw - 0.15) * SAFE, size, line_spacing=1.08)
              for p in points]
        need = max([d] + hs)
        return d, tx, tw, lh, off, hs, need

    # 丸は中（m）→ 小（s）。文字は本文の大きさから小さめの本文まで。それでも入らなければ縮める
    plans = [(k, v) for k in ("m", "s") for v in sorted({z["body"], z["body_small"]}, reverse=True)]
    v = z["body_small"] - 0.5
    while v >= 9.5:
        plans.append(("s", v))
        v -= 0.5
    for key, size in plans:
        d, tx, tw, lh, off, hs, need = layout(key, size)
        if n * (need + sep) - sep <= avail:
            break
    else:
        K.WARN.add("読み取りのポイントが枠に収まりません。文を短くするか、数を減らしてください")
    pitch = max(need + sep, min(avail / n, need + 1.0))
    if pitch >= avail / n - 1e-6:
        bottom = a.bottom                 # 行が面をちょうど等分している
    else:
        bottom = min(a.bottom, top + pitch * (n - 1) + need + 0.3)
    style.panel(s, Area(a.x, a.y, a.w, bottom - a.y))
    style.sub_heading(s, a.x + pad, head_y, a.w - 2 * pad, title, z["item_head"])
    y = top
    for i, (p, h) in enumerate(zip(points, hs)):
        K.max_chars_warn(p, {1: 40, 2: 40, 3: 25}.get(len(points), 16), "読み取りのポイント")
        style.marker(s, a.x + pad, y + max(0.0, (lh - d) / 2), i + 1, key)
        style.text(s, tx, y + off, _balanced_w(p, tw - 0.15, size) + 0.15, max(lh, h - off) + 0.04,
                   p, size, c["text"], anchor="t", fit_min=9.5, label="読み取りのポイント",
                   line_spacing=1.08)
        if i < n - 1:
            style.row_rule(s, a.x + pad, a.right - pad, y + need + (pitch - need) / 2)
        y += pitch


def _fits(y_end, a, what):
    """一覧の下端が本文の枠に収まっているか確かめ、はみ出すなら警告する。"""
    if y_end > a.bottom + 0.02:
        K.WARN.add(f"{what}が枠に収まりません。項目を減らすか、callout・note・lead を外してください")


# ====================================================================== グラフ

_CHART = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
          "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE,
          "stacked": XL_CHART_TYPE.COLUMN_STACKED}


_KIND_JA = {"column": "縦棒グラフ", "bar": "横棒グラフ", "line": "折れ線グラフ", "pie": "円グラフ",
            "stacked": "積み上げ縦棒グラフ"}


def _chart_alt(kind, title, cats, data, unit):
    """読み上げ用の代替テキスト（Office 基準 S3）。種類・見出し・値を短くまとめる。"""
    u = unit or ""
    parts = []
    for name, nums in data:
        vals = "、".join(f"{c} —" if v is None else f"{c} {v:g}{u}" for c, v in zip(cats, nums))
        parts.append(f"{name}：{vals}" if name else vals)
    text = f"{_KIND_JA.get(kind, 'グラフ')}。{K.plain(K.as_text(title))}。" + "／".join(parts)
    return text if len(text) <= 300 else text[:299] + "…"


def _font(font, style, size=None, color=None, bold=None):
    """グラフの文字の大きさ・色・書体（日本語の書体も）を指定する。"""
    if size is not None:
        font.size = Pt(max(size, 14))   # グラフの目盛・凡例・数値は 14pt 以上（Office 基準 S9）
    if color is not None:
        font.color.rgb = K.rgb(color)
    if bold is not None:
        font.bold = bold
    font.name = style.font("latin")
    rPr = font._rPr
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.Element(qn("a:ea"))
        rPr.find(qn("a:latin")).addnext(ea)
    ea.set("typeface", style.font("ja"))


def _labels(dl, style, size, color, fmt, position=None, bold=False):
    """データラベル（値）の書式。"""
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.show_value = True
    if position is not None:
        dl.position = position
    _font(dl.font, style, size, color, bold)


def _point_label(pt, style, size, color, fmt, position, pie=False):
    """1つの点（棒・扇形）のラベルだけを太字で目立たせる（書式・表示する内容は系列と同じ）。"""
    dl = pt.data_label
    _font(dl.font, style, size, color, True)
    dl.position = position
    d = dl._dLbl
    old = d.find(qn("c:numFmt"))
    if old is not None:
        d.remove(old)
    num = etree.Element(qn("c:numFmt"), formatCode=fmt, sourceLinked="0")
    for tag in ("c:spPr", "c:txPr", "c:dLblPos", "c:showLegendKey"):
        el = d.find(qn(tag))
        if el is not None:
            el.addprevious(num)
            break
    if pie:
        for tag, val in (("c:showVal", "0"), ("c:showCatName", "1"), ("c:showPercent", "1")):
            d.find(qn(tag)).set("val", val)
        _separator(d, "\n")


def _series_fills(style, n, hl):
    """系列（円グラフは項目）の色。強調があれば、その系列だけ主色にし、ほかは淡い灰色にする。
    淡い灰色は、淡い文字の色と罫線の色を混ぜて作る（暖かい配色のテーマでは暖かい灰色になる）。"""
    c = style.c
    base = style.series_colors()
    if hl is None:
        # 色が足りなければ、同じ色を淡くして続ける
        return [base[i % len(base)] if i < len(base) else _mix(base[i % len(base)], WHITE, 0.45)
                for i in range(n)]
    muted = [_mix(_faint(c), _grid(c), t) for t in (0.5, 0.75, 0.95, 0.62, 0.85, 1.0)]
    out = []
    k = 0
    for i in range(n):
        if i == hl:
            out.append(base[0])
        else:
            out.append(muted[k % len(muted)])
            k += 1
    return out


def _no_line(fmt):
    fmt.line.fill.background()


def _separator(dLbls, text):
    """データラベルの区切り（円グラフで「項目名」と「%」を改行で分ける）。"""
    old = dLbls.find(qn("c:separator"))
    if old is not None:
        dLbls.remove(old)
    sep = etree.Element(qn("c:separator"))
    sep.text = text
    after = dLbls.find(qn("c:showBubbleSize"))
    if after is None:
        after = dLbls.find(qn("c:showPercent"))
    if after is not None:
        after.addnext(sep)
    else:
        dLbls.append(sep)


def _nice_axis(lo, hi, n_max=5):
    """値の軸の最小・最大・目盛の間隔。間隔は切りのよい数（1・2・2.5・5 の 10 倍ずつ）で、
    目盛線が n_max + 1 本以下になるものを選ぶ（自動だと 0〜500 を 50 ずつ 11 本などになる）。"""
    if hi - lo <= 0:
        hi = lo + (abs(lo) or 1.0)
    raw = (hi - lo) / n_max
    mag = 10 ** math.floor(math.log10(raw))
    for m in (1, 2, 2.5, 5, 10, 20, 25, 50):
        step = m * mag
        if step < raw * (1 - 1e-9):
            continue
        vmin = math.floor(lo / step + 1e-9) * step
        vmax = math.ceil(hi / step - 1e-9) * step
        if (vmax - vmin) / step <= n_max + 1e-9:
            return vmin, vmax, step
    return lo, hi, (hi - lo) / n_max


def _step_fmt(step):
    """目盛の数の書式（間隔が整数なら整数、0.5 なら小数1桁、0.25 なら2桁）。"""
    t = f"{step:.6f}".rstrip("0").rstrip(".")
    dec = len(t.split(".")[1]) if "." in t else 0
    return "#,##0" + ("." + "0" * min(dec, 3) if dec else "")


def _as_list(v):
    """値のリストにそろえる（1つだけの値も受け付ける）。"""
    if v is None:
        return []
    if isinstance(v, (str, int, float, dict)):
        return [v]
    return list(v)


@register("chart")
def chart(style, sd):
    """グラフ（PowerPoint で編集できるグラフ）。

    chart: column（縦棒）/ bar（横棒）/ line（折れ線）/ pie（円）/ stacked（積み上げ縦棒）、
    categories: 項目、series: [{name, values}]、unit: 単位、
    highlight: 強調する系列の番号（円グラフと、系列が1つのグラフは項目の番号）、
    points: 右側の読み取りのポイント（任意、4個まで）、points_title: その見出し。
    強調があると、その系列（項目）だけ主色にして値を添え、ほかは淡い灰色にする。
    値の軸の目盛は切りのよい間隔で 6 本まで。
    """
    s, a = style.body(sd)
    c = style.c
    z = _z(style)
    kind = str(sd.get("chart", "column")).strip().lower()
    if kind not in _CHART:
        K.WARN.add(f"chart は column / bar / line / pie / stacked のどれかにしてください"
                   f"（いまは {kind}）。縦棒で描きます")
        kind = "column"
    cats = [K.plain(K.as_text(x)) for x in _as_list(sd.get("categories"))]
    series = [dict(x) for x in _as_list(sd.get("series")) if isinstance(x, dict)]
    if not cats or not series:
        K.WARN.add("グラフには categories と series が必要です")
        return
    if kind == "pie" and len(series) > 1:
        K.WARN.add("円グラフは系列を1つだけ描きます（2つ目以降は使いません）")
        series = series[:1]
    if kind == "stacked" and len(series) == 1:
        # 系列が1つの積み上げは縦棒と同じ形（値は棒の上に出す）
        kind = "column"
    if len(series) > 5:
        K.WARN.add(f"グラフの系列が {len(series)} 個あります（5 個以内を推奨）")
    limit = 6 if kind == "pie" else 12
    if len(cats) > limit:
        K.WARN.add(f"グラフの項目が {len(cats)} 個あります（{limit} 個以内を推奨）")
    data = []
    raw_all = []
    for sr in series:
        vals = _as_list(sr.get("values"))
        name = K.plain(K.as_text(sr.get("name", "")))
        if len(vals) != len(cats):
            K.WARN.add(f"系列「{name}」の値が {len(vals)} 個です（categories は {len(cats)} 個）")
            vals = (vals + [None] * len(cats))[:len(cats)]
        nums = []
        for v in vals:
            n = _num(v)
            if n is None and v not in (None, ""):
                K.WARN.add(f"系列「{name}」の値「{v}」が数ではありません")
            nums.append(n)
        raw_all += vals
        data.append((name, nums))
    hl = sd.get("highlight")
    # 系列が1つなら、強調は項目（棒・点）の番号として扱う
    by_cat = kind == "pie" or len(series) == 1
    n_hl = len(cats) if by_cat else len(series)
    if hl is not None:
        try:
            hl = int(hl)
        except (TypeError, ValueError):
            hl = -1
        if not 0 <= hl < n_hl:
            K.WARN.add(f"highlight は 0〜{n_hl - 1} の番号にしてください（いまは {sd.get('highlight')}）")
            hl = None
    fmt = _fmt(_decimals(raw_all))

    # ---- 配置: グラフはカードに入れ、ポイントがあれば右に置く
    # （右の欄は3列の段組みの右の1列と同じ幅。ほかの型の3列と端がそろう）
    points = _texts(sd.get("points"))
    if len(points) > 4:
        K.WARN.add(f"読み取りのポイントが {len(points)} 個あります（4 個以内を推奨）")
    gap = style.gutter()
    pw = (a.w - 2 * gap) / 3 if points else 0.0
    box = Area(a.x, a.y, a.w - (pw + gap if points else 0.0), a.h)
    style.plain_card(s, box)
    pad = 0.22
    unit = K.as_text(sd.get("unit", "") or "")
    top = box.y + pad
    if unit and kind != "pie":
        style.text(s, box.x + pad + 0.04, top - 0.04, 3.0, 0.3, f"（単位：{unit}）", z["caption"],
                   c["text_muted"], anchor="m", fit_min=9, label="単位")
        top += 0.26
    fx, fw = box.x + pad, box.w - 2 * pad
    fh = box.bottom - pad - top

    cd = CategoryChartData(number_format=fmt)
    cd.categories = cats
    for name, nums in data:
        cd.add_series(name, nums)
    gf = s.shapes.add_chart(_CHART[kind], Inches(fx), Inches(top), Inches(fw), Inches(fh), cd)
    gf._element.nvGraphicFramePr.cNvPr.set("descr", _chart_alt(kind, sd.get("title", ""), cats, data, unit))
    ch = gf.chart
    ch.has_title = False
    _font(ch.font, style, z["body_small"], c["text"])
    plot = ch.plots[0]
    many = len(cats) > 8
    # 値のラベル（多いときは注釈の大きさ）と、強調する値のラベル（一段大きく）
    lab_size = z["caption"] if many else z["body_small"]
    lab_big = z["body_small"] if many else z["body"]

    if kind == "pie":
        fills = _series_fills(style, len(cats), hl)
        ch.has_legend = False
        ser = plot.series[0]
        for i, col in enumerate(fills):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = K.rgb(col)
            pt.format.line.color.rgb = K.rgb(WHITE)
            pt.format.line.width = Pt(1.5)
        # ラベルは系列に付ける（強調する扇形だけを上書きできるように）
        dl = ser.data_labels
        dl.number_format = "0%"
        dl.number_format_is_linked = False
        dl.show_value = False
        dl.show_category_name = True
        dl.show_percentage = True
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        _font(dl.font, style, z["body_small"], c["text"])
        _separator(dl._element, "\n")
        if hl is not None:
            _point_label(ser.points[hl], style, z["body"], c["heading"], "0%",
                         XL_LABEL_POSITION.OUTSIDE_END, pie=True)
        _chart_points(style, s, a, box, gap, pw, points, sd)
        return

    hl_ser = None if by_cat else hl
    hl_cat = hl if by_cat else None
    fills = _series_fills(style, len(series), hl_ser)
    strong = fills[0]
    if hl_cat is not None and kind != "line":
        # 系列が1つで項目を強調するときは、棒を淡くして、その項目だけ主色にする
        # （折れ線は線の色はそのままで、その点の印を大きくする）
        fills = _series_fills(style, 2, 1)[:1]
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        _font(ch.legend.font, style, z["body_small"], c["text"])
    va, ca = ch.value_axis, ch.category_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = K.rgb(_grid(c))
    va.major_gridlines.format.line.width = Pt(0.75)
    _no_line(va.format)
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    _font(va.tick_labels.font, style, z["caption"], c["text_muted"])
    ca.format.line.color.rgb = K.rgb(_faint(c))
    ca.format.line.width = Pt(1.0)
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    nums_all = [v for _, ns in data for v in ns if v is not None]
    if nums_all and min(nums_all) < 0:
        # 負の値があると項目名が 0 の線に重なるので、軸の端（縦棒は下、横棒は左）に置く
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    _font(ca.tick_labels.font, style, z["caption"] if many else z["body_small"], c["text"])
    if kind == "bar":
        # 上から categories の順に並べ、値の目盛は下に置く
        ca.reverse_order = True
        crosses = va._element.find(qn("c:crosses"))
        if crosses is not None:
            crosses.set("val", "max")
    if kind in ("column", "bar", "stacked"):
        plot.gap_width = 55 if len(series) > 1 else 70
        if kind != "stacked":
            plot.overlap = -8 if len(series) > 1 else 0
    single = len(series) == 1
    shows = []
    for i, (ser, col) in enumerate(zip(plot.series, fills)):
        emph = i == hl_ser
        # 値のラベルは、強調する系列か、数が少ないときだけ付ける（多いと重なって読めない）
        show = emph or (hl_ser is None and (len(series) * len(cats) <= 16
                                            or (single and len(cats) <= 12)))
        shows.append(show)
        if kind == "line":
            ser.smooth = False
            ser.format.line.color.rgb = K.rgb(col)
            ser.format.line.width = Pt(3.0 if emph else 2.25)
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 8 if emph else 7
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = K.rgb(col)
            ser.marker.format.line.color.rgb = K.rgb(WHITE)
            if show and (single or emph):
                _labels(ser.data_labels, style, lab_size, c["heading"] if emph else c["text"],
                        fmt, XL_LABEL_POSITION.ABOVE, bold=emph)
        else:
            ser.invert_if_negative = False
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = K.rgb(col)
            _no_line(ser.format)
            if kind == "stacked":
                if len(cats) <= 10:
                    _labels(ser.data_labels, style, z["caption"], _on(c, col), fmt,
                            XL_LABEL_POSITION.CENTER)
            elif show:
                _labels(ser.data_labels, style, lab_size, c["heading"] if emph else c["text"],
                        fmt, XL_LABEL_POSITION.OUTSIDE_END, bold=emph)
        if hl_cat is not None:
            pt = ser.points[hl_cat]
            if kind == "line":
                pt.marker.style = XL_MARKER_STYLE.CIRCLE
                pt.marker.size = 12
                pt.marker.format.fill.solid()
                pt.marker.format.fill.fore_color.rgb = K.rgb(WHITE)
                pt.marker.format.line.color.rgb = K.rgb(strong)
                pt.marker.format.line.width = Pt(2.75)
            else:
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = K.rgb(strong)
            if show:
                _point_label(pt, style, lab_big, c["heading"], fmt,
                             XL_LABEL_POSITION.ABOVE if kind == "line"
                             else XL_LABEL_POSITION.OUTSIDE_END)

    # ---- 値の軸: 切りのよい間隔で、目盛線は 6 本まで
    if nums_all:
        if kind == "stacked":
            cols = list(zip(*(ns for _, ns in data)))
            hi = max([0.0] + [sum(v for v in col if v and v > 0) for col in cols])
            lo = min([0.0] + [sum(v for v in col if v and v < 0) for col in cols])
        elif kind == "line":
            lo, hi = min(nums_all), max(nums_all)
            span = (hi - lo) or abs(hi) or 1.0
            lo2 = lo - span * 0.1
            lo = 0.0 if lo >= 0 > lo2 else lo2
            hi = hi + span * (0.12 if any(shows) else 0.05)
        else:
            lo, hi = min([0.0] + nums_all), max([0.0] + nums_all)
            if any(shows) or hl_cat is not None:
                # 棒の外に出す値のラベルが枠に当たらないよう、少しだけ上（下）をあける
                hi, lo = hi * 1.08, lo * 1.08
        vmin, vmax, step = _nice_axis(lo, hi)
        va.minimum_scale = vmin
        va.maximum_scale = vmax
        va.major_unit = step
        va.tick_labels.number_format = _step_fmt(step)
        va.tick_labels.number_format_is_linked = False
    _chart_points(style, s, a, box, gap, pw, points, sd)


def _chart_points(style, s, a, box, gap, pw, points, sd):
    """グラフの右に読み取りのポイントを置く（あるときだけ）。"""
    if points:
        _points_panel(style, s, Area(box.right + gap, a.y, pw, a.h), points,
                      K.as_text(sd.get("points_title") or "読み取りのポイント"))


# ====================================================================== スケジュール

def _period_index(v, periods, default, what):
    """start・end を periods の位置にする（数字のほか、期間の名前でもよい）。"""
    n = len(periods)
    if v is None or v == "":
        return default
    if isinstance(v, str) and v in periods:
        return periods.index(v)
    try:
        i = int(v)
    except (TypeError, ValueError):
        K.WARN.add(f"{what}「{v}」は periods の位置（0〜{n - 1}）か期間の名前にしてください")
        return default
    if not 0 <= i < n:
        K.WARN.add(f"{what} {i} が periods の範囲（0〜{n - 1}）の外です")
        i = max(0, min(n - 1, i))
    return i


@register("schedule")
def schedule(style, sd):
    """スケジュール（ガントチャート）。periods: 期間の見出し、
    tasks: [{name, start, end, owner, milestone, label}]（8個まで）。
    start・end は periods の位置（0 から）か期間の名前。label は帯（節目は印の横）に添える短い文字。
    文字の大きさと見出しの行は表の型と同じ（表の仲間としてそろえる）。担当のない行は「—」。"""
    s, a = style.body(sd)
    c = style.c
    z = _z(style)
    periods = [K.plain(K.as_text(p)) for p in _as_list(sd.get("periods"))]
    tasks = _rename(as_items(sd.get("tasks")), "name")
    if not periods or not tasks:
        K.WARN.add("スケジュールには periods と tasks が必要です")
        return
    check_count(tasks, 1, 8, "スケジュールのタスク")
    if len(periods) > 12:
        K.WARN.add(f"期間が {len(periods)} 個あります（12 個以内を推奨。月を四半期にまとめるなど）")
    tc = style.table_colors()
    np_ = len(periods)
    n = len(tasks)
    hs, ts = _table_sizes(style, z)
    hh = max(0.38, _line_h(hs) + 0.14)       # 表の見出しの行と同じ高さ
    # 行の高さは本文の枠から決める（少ないときは高く、多いときは詰める）
    rh = min(1.0 if n <= 3 else 0.74, (a.h - hh) / n)
    name_size = own_size = ts
    if rh < 0.42:
        name_size = own_size = min(ts, z["caption"])
    if rh < 0.34:
        # 上限を超える数のときも、文字を小さくして枠に収める（0.28 インチまで）
        name_size = own_size = min(z["caption"], 10)
        rh = max(rh, 0.28)
    has_owner = any(t.get("owner") not in (None, "") for t in tasks)
    longest = max(K.text_em(K.plain(K.as_text(t.get("name", ""))), True) for t in tasks)
    nw = min(3.3, max(2.2, longest * name_size / 72 + 0.4))
    ow = 0.0
    if has_owner:
        lo = max(K.text_em(K.plain(K.as_text(t.get("owner", "")))) for t in tasks)
        ow = min(2.0, max(1.2, lo * own_size / 72 + 0.34))
    tx = a.x + nw + ow
    pw = (a.right - tx) / np_
    y0 = a.y
    body_h = rh * n
    pad = 0.14
    _fits(y0 + hh + body_h, a, "スケジュール")

    # ---- 見出しの行（表と同じ色・大きさ。寄せは列の中身にそろえる）
    K.add_box(s.shapes, a.x, y0, a.w, hh, fill=tc["head_bg"])
    style.text(s, a.x + pad, y0, nw - 2 * pad, hh, "タスク", hs, tc["head_text"], bold=True,
               anchor="m", fit_min=9)
    if has_owner:
        style.text(s, a.x + nw + 0.04, y0, ow - 0.08, hh, "担当", hs, tc["head_text"], bold=True,
                   align="c", anchor="m", fit_min=9)
    for j, p in enumerate(periods):
        style.text(s, tx + j * pw + 0.03, y0, pw - 0.06, hh, p, hs, tc["head_text"], bold=True,
                   align="c", anchor="m", fit_min=9, label="期間の見出し")

    # ---- 行の地と罫線
    yb = y0 + hh
    for i in range(n):
        y = yb + i * rh
        bg = tc["row_alt"] if i % 2 else tc["row_bg"]
        K.add_box(s.shapes, a.x, y, a.w, rh, fill=bg)
    for i in range(n):
        y = yb + (i + 1) * rh
        K.add_line(s.shapes, a.x, y, a.right, y, tc["line"], 0.75)
    for j in range(1, np_):
        x = tx + j * pw
        K.add_line(s.shapes, x, yb, x, yb + body_h, tc["line"], 0.75, dash=None)
    for x in ([a.x + nw] + ([tx] if has_owner else [])):
        K.add_line(s.shapes, x, yb, x, yb + body_h, tc["line"], 0.75)
    for x in (a.x, a.right):
        K.add_line(s.shapes, x, yb, x, yb + body_h, tc["line"], 0.75)

    # ---- タスク名・担当・期間の帯
    bar = c["primary"]
    # 節目の印は強調色（帯の色と区別しやすい）
    mark = c["emph"]
    bh = min(0.32, rh * 0.5)
    cap = z["caption"]
    for i, t in enumerate(tasks):
        y = yb + i * rh
        name = K.as_text(t.get("name", ""))
        K.max_chars_warn(name, 20, "タスク名")
        style.text(s, a.x + pad, y, nw - 2 * pad, rh, name, name_size, tc["first_col"], bold=True,
                   anchor="m", fit_min=9.5, label="タスク名", line_spacing=0.95)
        if has_owner:
            if t.get("owner") not in (None, ""):
                style.text(s, a.x + nw + 0.06, y, ow - 0.12, rh, K.as_text(t["owner"]), own_size,
                           tc["text"], align="c", anchor="m", fit_min=9, label="担当",
                           line_spacing=0.95)
            else:
                _blank(style, s, a.x + nw + 0.06, y, ow - 0.12, rh, own_size)
        st = _period_index(t.get("start"), periods, 0, f"「{K.plain(name)}」の start")
        en = _period_index(t.get("end"), periods, st, f"「{K.plain(name)}」の end")
        if en < st:
            K.WARN.add(f"「{K.plain(name)}」の end が start より前です")
            st, en = en, st
        cy = y + rh / 2
        label = K.as_text(t.get("label", "") or "")
        if t.get("milestone"):
            d = min(0.38, rh * 0.62)
            mx = tx + (en + 0.5) * pw - d / 2
            K.add_box(s.shapes, mx, cy - d / 2, d, d, fill=mark, line=WHITE, line_w=1.25,
                      geom="diamond")
            if label:
                # 節目の補足（任意）は印の右に。右に余裕がなければ左に置く
                lw = min(2.4, K.text_em(K.plain(label)) * cap / 72 + 0.1)
                right = a.right - (mx + d + 0.06) >= lw
                lx = mx + d + 0.06 if right else max(tx, mx - 0.06 - lw)
                style.text(s, lx, y, min(lw, a.right - lx), rh, label, cap, c["text"],
                           align="l" if right else "r", anchor="m", fit_min=9, label="節目の補足")
        else:
            bx = tx + st * pw + 0.08
            bw = (en - st + 1) * pw - 0.16
            K.add_box(s.shapes, bx, cy - bh / 2, bw, bh, fill=bar, geom="round",
                      radius=bh * 0.35)
            if label:
                style.text(s, bx, cy - bh / 2, bw, bh, label, cap, _on(c, bar),
                           bold=True, align="c", anchor="m", fit_min=9, label="帯の文字")


# ====================================================================== 体制

def _leader(sd):
    """責任者（leader）を辞書にそろえる。文字列なら名前として扱う。

    lead は本文ページ共通の「リード文」と同じ名前なので、辞書のときだけ責任者として扱う。"""
    lead = sd.get("leader")
    if isinstance(lead, str) and lead.strip():
        lead = {"name": lead}
    if not isinstance(lead, dict):
        lead = sd.get("lead") if isinstance(sd.get("lead"), dict) else {}
    if lead and not (lead.get("name") or lead.get("role")):
        K.WARN.add("責任者（leader）には name か role を書いてください")
        lead = {}
    return lead


@register("team")
def team(style, sd):
    """体制図。leader: {role, name, note}（責任者。名前だけの文字列でもよい）、
    groups: [{role, name, note}]（2〜5個。name はメンバーのリストでもよい）。

    責任者は lead に辞書で書いてもよい（文字列の lead はリード文）。
    責任者の箱を a.y に置き、チームのカードは a.x〜a.right を style.gutter() の間隔で等分する。"""
    lead = _leader(sd)
    if isinstance(sd.get("lead"), dict):
        sd = {k: v for k, v in sd.items() if k != "lead"}
    s, a = style.body(sd)
    c = style.c
    z = _z(style)
    groups = _rename(as_items(sd.get("groups")), "role")
    check_count(groups, 2, 5, "体制のチーム")
    if not groups:
        return
    for g in groups:
        # 名前はリスト（メンバーの名前）でもよい
        names = g.get("name", g.get("members", ""))
        g["name"] = ("・".join(K.as_text(v) for v in names) if isinstance(names, list)
                     else K.as_text(names))
        g["role"] = K.as_text(g.get("role", ""))
    n = len(groups)
    gap = style.gutter()
    gw = (a.w - gap * (n - 1)) / n
    gx0 = a.x

    # ---- チームのカードの高さ（いちばん多いものに合わせる）
    inner_w = gw - 0.4
    for g in groups:
        K.max_chars_warn(g.get("role", ""), 14, "チーム名")
        if g.get("note"):
            K.max_chars_warn(g["note"], 40, "チームの補足")
    # チーム名の帯（kpi_tile の帯）の高さ
    head_h = 0.44

    def text_hs(ns, ts):
        """名前（人数）と補足の高さ（いちばん多いチームに合わせる）。"""
        name_h = max(K.text_height(K.plain(g["name"]), inner_w * SAFE, ns, True) for g in groups)
        note_h = max([K.text_height(K.plain(K.as_text(g["note"])), inner_w * SAFE, ts,
                                    line_spacing=1.05) for g in groups if g.get("note")] or [0.0])
        return max(0.34, name_h), note_h

    def need(icon, hs):
        name_h, note_h = hs
        return (head_h + 0.2 + (icon + 0.1 if icon else 0.0) + name_h
                + (0.2 + note_h + 0.08 if note_h else 0.0) + 0.22)

    # 責任者の箱は、書いてある行（役割・名前・補足）の高さ＋上下の余白
    lines = [(key, _line_h(size) + extra, size)
             for key, size, extra in (("role", z["caption"], 0.08), ("name", z["item_head"], 0.12),
                                      ("note", z["caption"], 0.06)) if lead.get(key)]
    text_h = sum(h for _, h, _ in lines)
    lead_h = max(1.0, text_h + 0.26) if lead else 0.0
    # 名前は項目の見出し（4チーム以上はゆったりした本文）、補足は本文（4チーム以上は小さめ）。
    # 高さが足りないときは、線の長さ → 人のアイコン → 文字の大きさ の順に詰める
    first = (z["item_head"], z["body"]) if n <= 3 else (z["body_large"], z["body_small"])
    plans = [first, (z["body"], z["body_small"]), (z["body_small"], z["caption"])]
    for name_size, note_size in plans:
        hs = text_hs(name_size, note_size)
        fit = next(((icon, conn) for icon in (0.5, 0.4, 0.0)
                    for conn in ((0.62, 0.36) if lead else (0.0,))
                    if need(icon, hs) <= a.h - lead_h - conn + 0.04), None)
        if fit:
            break
    else:
        fit = (0.0, 0.36 if lead else 0.0)
        K.WARN.add("体制図が枠に収まりません。補足を短くするか、callout・note・lead を外してください")
    icon, conn = fit
    name_h, note_h = hs
    card_h = min(need(icon, hs), a.h - lead_h - conn)
    y = a.y
    line = _faint(c)

    # ---- 責任者（カードの部品を濃い色で塗る。角丸・影はデザインに任せる）
    if lead:
        lw = min(4.8, max(3.4, a.w * 0.34))
        lx = a.x + (a.w - lw) / 2
        fill, fg = _strong(c)
        _recolor(style.plain_card(s, Area(lx, y, lw, lead_h)), fill)
        d = 0.56
        _person(s, lx + 0.28, y + (lead_h - d) / 2, d, WHITE, fill)
        txx = lx + 0.28 + d + 0.2
        tw = lx + lw - 0.24 - txx
        ty = y + (lead_h - text_h) / 2
        for key, h, size in lines:
            style.text(s, txx, ty, tw, h, K.as_text(lead[key]), size, fg, bold=key != "note",
                       anchor="m", fit_min=min(size, 12) if key == "name" else 9,
                       label={"role": "責任者の役割", "name": "責任者の名前",
                              "note": "責任者の補足"}[key])
            ty += h
        # 責任者 → 各チームへの線（白い地と 3:1 ほどの差のある色。淡すぎると見えない）
        cx = a.x + a.w / 2
        by = y + lead_h + conn * 0.5
        K.add_line(s.shapes, cx, y + lead_h, cx, by, line, 1.5)
        centers = [gx0 + i * (gw + gap) + gw / 2 for i in range(n)]
        K.add_line(s.shapes, centers[0], by, centers[-1], by, line, 1.5)
        for x in centers:
            K.add_line(s.shapes, x, by, x, y + lead_h + conn, line, 1.5)
        y += lead_h + conn

    # ---- チーム
    for i, g in enumerate(groups):
        x = gx0 + i * (gw + gap)
        box = Area(x, y, gw, card_h)
        inner = style.kpi_tile(s, box, {"label": g.get("role", "")}, i)
        iy = inner.y + 0.2
        if icon:
            _person(s, x + (gw - icon) / 2, iy, icon, c["bg_alt"], c["primary"])
            iy += icon + 0.1
        # 名前と補足は中央ぞろえ。折り返すときは各行の長さをそろえる（1文字だけ次の行に送らない）
        bw = _balanced_w(g["name"], inner_w, name_size, True)
        style.text(s, x + (gw - bw) / 2, iy, bw, name_h, g["name"], name_size, c["heading"],
                   bold=True, align="c", anchor="m", fit_min=10, label="チームの人数・名前")
        iy += name_h
        if g.get("note"):
            iy += 0.1
            K.add_line(s.shapes, x + gw * 0.3, iy, x + gw * 0.7, iy, _grid(c), 1.0)
            iy += 0.1
            bw = _balanced_w(g["note"], inner_w, note_size)
            style.text(s, x + (gw - bw) / 2, iy, bw, max(0.2, y + card_h - iy - 0.12),
                       K.as_text(g["note"]), note_size, c["text"], align="c", anchor="t",
                       fit_min=9.5, label="チームの補足", line_spacing=1.05)


# ====================================================================== 費用

def _totals(v):
    """total を [{label, amount, note}] のリストにそろえる（数字だけなら「合計」の金額）。"""
    if v is None or v == "" or v == []:
        return []
    if isinstance(v, dict) or not isinstance(v, list):
        v = [v]
    out = []
    for t in v:
        if isinstance(t, (str, int, float)) and not isinstance(t, bool):
            t = {"label": "合計", "amount": t}
        if not isinstance(t, dict):
            continue
        if t.get("amount") in (None, ""):
            K.WARN.add(f"合計「{t.get('label', '')}」に amount（金額）がありません（表示しません）")
            continue
        out.append(t)
    if len(out) > 2:
        K.WARN.add(f"合計（total）が {len(out)} 個あります（2 個まで表示します）")
        out = out[:2]
    return out


def _table_h(header, rows, ratios, w, size):
    """render_table と同じ見積もりで、表の高さ（インチ）を返す。"""
    widths = [w * r / float(sum(ratios)) for r in ratios]
    total = 0.0
    for i, r in enumerate([header] + rows):
        lines = max(K.count_lines(K.plain(str(t)), cw - 0.24, size, i == 0)
                    for t, cw in zip(r, widths))
        total += max(lines * size * K.LINE_HEIGHT / 72 + 0.14, 0.38)
    return total


@register("cost")
def cost(style, sd):
    """費用・見積。items: [{name, detail, amount}]（1〜8個）、unit: 単位、
    total: {label, amount, note}（初期費用と運用費のように2つならリスト。数字だけでもよい）、
    notes: 前提・注記。total がなく金額がすべて数なら、2項目以上のとき合計を自動で出す。
    右の列（合計）は3列の段組みの右の1列と同じ幅。表の文字は表の型と同じ大きさ。"""
    s, a = style.body(sd)
    c = style.c
    z = _z(style)
    items = _rename(as_items(sd.get("items")), "name")
    check_count(items, 1, 8, "費用の項目")
    if not items:
        return
    unit = K.as_text(sd.get("unit", "") or "")
    totals = _totals(sd.get("total"))
    amounts = [it.get("amount") for it in items]
    nums = [_num(v) for v in amounts]
    if all(v is not None for v in nums):
        dec = _decimals(amounts)
        sm = sum(nums)
        sm_txt = f"{round(sm, dec):,.{dec}f}"
        if not totals and len(items) >= 2:
            totals = [{"label": "合計", "amount": sm_txt}]
        elif totals:
            tv = _num(totals[0].get("amount"))
            if tv is not None and abs(tv - sm) > 1e-6:
                K.WARN.add(f"合計 {_amount(totals[0].get('amount'))} が内訳の和 {sm_txt} と"
                           "一致しません。数字を確認してください")
    notes = _texts(sd.get("notes"))
    has_detail = any(it.get("detail") not in (None, "") for it in items)
    nt = len(totals)
    _, ts = _table_sizes(style, z)

    gap = style.gutter()
    rw = (a.w - 2 * gap) / 3 if totals else 0.0
    tw = a.w - (rw + gap if rw else 0.0)
    head = ["項目"] + (["内容"] if has_detail else []) + [f"金額（{unit}）" if unit else "金額"]
    rows = []
    for it in items:
        K.max_chars_warn(it.get("name", ""), 20, "費用の項目名")
        row = [K.as_text(it.get("name", ""))]
        if has_detail:
            row.append(K.as_text(it.get("detail", "") or ""))
        row.append(_amount(it.get("amount", "")))
        rows.append(row)
    ratios = [2.2, 3.0, 1.5] if has_detail else [3.2, 1.6]

    # ---- 配置を決める
    # 合計のタイル1枚: 項目名・数字・補足が読める高さ（th_min）〜 大きく見せる高さ（th_max）
    th_max = 2.15 if nt == 1 else 2.05
    th_min = 1.4 + (0.25 if any(t.get("note") for t in totals) else 0.0)
    # 前提・注記は、右の列（合計の下）に入ればそこへ。入らなければ表の下に置く
    notes_right = bool(notes and totals) and (
        nt * th_min + 0.24 * (nt - 1) + _notes_h(style, notes, rw) + 0.24 <= a.h)
    # 表の合計の行（右の数字と同じ合計）は、入りきらないときは省く
    cands = [False]
    if nt and len(items) > 1:
        t0 = totals[0]
        trow = ([K.as_text(t0.get("label", "合計"))] + ([""] if has_detail else [])
                + [_amount(t0["amount"])])
        cands = [True, False]
    need = {tr: _table_h(head, rows + ([trow] if tr else []), ratios, tw, ts) for tr in cands}
    nh, compact = 0.0, False
    if notes and not notes_right:
        # 表の下: 面の形 → 1行にまとめた小さな文字 の順に、表と合わせて入るものを選ぶ
        found = None
        for cp in (False, True):
            for tr in cands:
                h = _notes_h(style, notes, tw, cp)
                if need[tr] + _notes_gap(cp) + h <= a.h:
                    found = (tr, cp, h)
                    break
            if found:
                break
        if found is None:
            tr = cands[-1]
            h = a.h - need[tr] - _notes_gap(True)
            if h >= K.text_height(K.plain(_notes_line(notes)), tw - 0.08, 8) + 0.04:
                K.WARN.add("前提・注記が枠に収まらないため、小さな文字で1つにまとめました。"
                           "注記を減らすか、短くしてください")
                found = (tr, True, h)
            else:
                K.WARN.add("前提・注記を置く場所がありません（表示しません）。項目か注記を減らしてください")
                found = (tr, False, 0.0)
        total_row, compact, nh = found
    else:
        total_row = next((tr for tr in cands if need[tr] <= a.h), cands[-1])
    if total_row:
        rows.append(trow)
    table_h = a.h - (nh + _notes_gap(compact) if nh else 0.0)
    last = len(rows) - 1
    amt = len(head) - 1
    tint = c.get("tint", c["bg_alt"])

    def cell(i, j, text):
        if total_row and i == last:
            return {"fill": tint, "color": c["heading"], "bold": True}
        return None

    aligns = {str(j): "l" for j in range(amt)}
    aligns[str(amt)] = "r"
    used = render_table(style, s, Area(a.x, a.y, tw, table_h), head, rows, ratios, aligns, None,
                        cell, size_max=ts, size_min=9.5, label="費用の表")
    if nh:
        ny = a.y + used + _notes_gap(compact)
        _notes(style, s, Area(a.x, ny, tw, min(nh, a.bottom - ny)), notes, compact)
    if not totals:
        return

    # ---- 右: 合計の数字と、前提・注記
    rx = a.x + tw + gap
    y = a.y
    notes_need = _notes_h(style, notes, rw) + 0.24 if notes_right else 0.0
    th = min(th_max, (a.h - notes_need - 0.24 * (nt - 1)) / nt)
    if notes_right and nt == 1:
        # 右の列（合計＋注記）の下端が、なるべく表の下端にそろうように合計の高さを決める
        th = min(th, max(th_min, used - notes_need))
    if th < th_min - 0.01:
        K.WARN.add("合計の数字を置く高さが足りません。合計を1つにするか、callout・note を外してください")
    # 数字の大きさは数字の強調（kpi）の段階から: 1つなら3枚並びと同じ、2つならその一段下
    big = _as_list(style.size("kpi_number", [72, 64, 56, 48])) or [56]
    big_size = big[min(2 if nt == 1 else 3, len(big) - 1)]
    for i, t in enumerate(totals):
        tile = Area(rx, y, rw, th)
        inner = style.kpi_tile(s, tile, {"label": K.as_text(t.get("label", "合計"))}, i)
        note = K.as_text(t.get("note", "") or "")
        note_h = 0.42 if note else 0.0
        box_h = inner.h - note_h - 0.14
        # 高さが足りないときは、1pt ずつ縮めずに数字の段階を1つずつ下げる
        steps = [v for v in big if v <= big_size] or [big_size]
        size = next((v for v in steps if v * K.LINE_HEIGHT / 72 <= box_h + 0.05), steps[-1])
        style.kpi_value(s, Area(inner.x + 0.15, inner.y + 0.04, inner.w - 0.3, box_h),
                        _amount(t["amount"]), unit, size)
        if note:
            K.max_chars_warn(note, 30, "合計の補足")
            style.kpi_note(s, Area(inner.x, inner.bottom - note_h - 0.1, inner.w, note_h), note)
        y += th + 0.24
    if notes_right:
        bottom = min(a.bottom, max(a.y + used, y + _notes_h(style, notes, rw)))
        _notes(style, s, Area(rx, y, rw, bottom - y), notes)


def _notes_line(notes):
    """前提・注記を1行にまとめた形（場所が足りないとき。表の下の脚注のように置く）。"""
    return "※ " + "　／　".join(notes)


def _notes_gap(compact):
    """表と前提・注記の間。"""
    return 0.12 if compact else 0.2


def _notes_h(style, notes, w, compact=False):
    """前提・注記に要る高さ。compact は面を使わず1行にまとめた形。"""
    z = _z(style)
    if compact:
        return K.text_height(K.plain(_notes_line(notes)), w - 0.08, z["caption"],
                             line_spacing=1.05) + 0.06
    head = max(0.36, _line_h(z["item_head"]) + 0.04)
    return 0.2 + head + 0.08 + K.text_height("\n".join(K.plain(n) for n in notes), w - 0.66,
                                             z["body_small"], line_spacing=1.05) \
        + 0.06 * len(notes) + 0.22


def _notes(style, s, a, notes, compact=False):
    """前提・注記の面（淡い色の面に小見出しと箇条書き）。compact なら1行にまとめた小さな文字。"""
    c = style.c
    z = _z(style)
    if compact:
        style.text(s, a.x + 0.04, a.y, a.w - 0.08, a.h, _notes_line(notes), z["caption"],
                   c["text_muted"], anchor="t", fit_min=8, label="費用の注記", line_spacing=1.05)
        return
    style.panel(s, a)
    y = a.y + 0.16
    y += style.sub_heading(s, a.x + 0.22, y, a.w - 0.44, "前提・注記", z["item_head"]) + 0.06
    style.text(s, a.x + 0.22, y, a.w - 0.44, max(0.2, a.bottom - y - 0.14), notes, z["body_small"],
               c["text"], anchor="t", fit_min=9, label="費用の注記", bullet="・",
               bullet_color=c["primary"], para_space=3, line_spacing=1.05)


# ====================================================================== リスク

@register("risks")
def risks(style, sd):
    """リスクと対策。items: [{risk, level, measure, owner}]（6個まで）。level は 高・中・低。

    列見出しを a.y に置き、その下に行（カード）を上から詰めて並べる（next-actions と同じ組み方）。
    「高」の行は強調の面に「高」の色の線を添える。色はテーマの levels だけを使う。"""
    s, a = style.body(sd)
    c = style.c
    z = _z(style)
    items = _rename(as_items(sd.get("items")), "risk")
    check_count(items, 1, 6, "リスク")
    if not items:
        return
    n = len(items)
    has_owner = any(it.get("owner") not in (None, "") for it in items)
    head_h = _head_h(z)
    avail = a.h - head_h - 0.08
    # 行の高さは本文の枠から決める。足りなければ間隔と文字を詰める
    gap = 0.2 if n <= 4 else 0.14
    rh = min(1.4 if n == 1 else 1.3 if n == 2 else 1.2, (avail - gap * (n - 1)) / n)
    # リスクは項目の見出し、対策は本文（3つまではゆったり）。対策は見出しより大きくしない
    if n <= 3:
        rsize, msize = z["item_head"], z["body_large"]
    elif n <= 4:
        rsize, msize = z["item_head"], z["body"]
    else:
        rsize, msize = z["body_large"], z["body"]
    inset = 0.06
    compact = rh < 0.62
    if compact:
        gap, inset = 0.08, 0.04
        rsize, msize = z["body"], z["body_small"]
        rh = (avail - gap * (n - 1)) / n
    if rh < 0.46:
        rh = 0.46
    msize = min(msize, rsize)
    total = head_h + 0.08 + rh * n + gap * (n - 1)
    _fits(a.y + total, a, "リスクの一覧")
    y = a.y
    pad = 0.24
    lvw = 0.8
    aw = 0.5
    ow = 1.8 if has_owner else 0.0
    flex = a.w - 2 * pad - lvw - 0.3 - aw - (ow + 0.2 if ow else 0.0)
    risk_w = flex * 0.44
    meas_w = flex - risk_w
    x_lv = a.x + pad
    x_risk = x_lv + lvw + 0.3
    x_arrow = x_risk + risk_w
    x_meas = x_arrow + aw
    x_own = x_meas + meas_w + 0.2
    _head(style, s, x_lv - 0.1, y, lvw + 0.2, head_h, "影響度", align="c")
    _head(style, s, x_risk, y, risk_w, head_h, "リスク")
    _head(style, s, x_meas, y, meas_w, head_h, "対策")
    if has_owner:
        _head(style, s, x_own, y, ow, head_h, "担当")
    y += head_h + 0.08
    high = style.level_style("高")
    chip_size = z["caption"] if compact else z["body_small"]
    ph = min(_pill_h(chip_size), rh - 0.1)
    ah = min(0.3, rh - 0.16)
    cap = z["caption"]
    for i, it in enumerate(items):
        K.max_chars_warn(it.get("risk", ""), 30, "リスク")
        K.max_chars_warn(it.get("measure", ""), 40, "対策")
        level = K.as_text(it.get("level", "") or "").strip()
        if level and level not in ("高", "中", "低", "大", "小", "high", "mid", "medium", "low"):
            K.WARN.add(f"リスクの level「{level}」は 高・中・低 のどれかにしてください")
        row = Area(a.x, y, a.w, rh)
        if level and style.level_style(level) == high:
            # 「高」と同じ色になる項目（高・大 など）は、強調の面に「高」の色の線
            bg = _attention(style, s, row, high[0])
        else:
            bg = _fill_of(style.plain_card(s, row), c["card_bg"])
        if level:
            fill, fg = style.level_style(level)
            _chip(style, s, x_lv + 0.04, y + (rh - ph) / 2, lvw - 0.08, ph, level, fill, fg,
                  chip_size)
        style.text(s, x_risk, y + inset, risk_w - 0.1, rh - 2 * inset,
                   K.as_text(it.get("risk", "")), rsize, c["heading"], bold=True, anchor="m",
                   fit_min=10, label="リスク", line_spacing=1.05)
        style.arrow(s, x_arrow + (aw - 0.2) / 2, y + (rh - ah) / 2, 0.2, ah, "right",
                    color=_faint(c))
        style.text(s, x_meas + 0.06, y + inset, meas_w - 0.12, rh - 2 * inset,
                   K.as_text(it.get("measure", "")), msize, c["text"], anchor="m", fit_min=9.5,
                   label="対策", line_spacing=1.05)
        if has_owner:
            if it.get("owner") not in (None, ""):
                style.text(s, x_own, y + inset, ow - 0.1, rh - 2 * inset, K.as_text(it["owner"]),
                           cap, _ink(c, c["text_muted"], bg), anchor="m", fit_min=9,
                           label="担当")
            else:
                _blank(style, s, x_own, y + inset, ow - 0.1, rh - 2 * inset, cap, align="l")
        y += rh + gap


# ====================================================================== 課題・進捗

_STATUS_ORDER = ("完了", "進行中", "遅延", "未着手")


@register("status")
def status(style, sd):
    """課題・進捗の一覧。items: [{name, status, progress, owner, due, note}]（8個まで）。
    status は 完了・進行中・遅延・未着手、progress は 0〜100（任意。'50%' でもよい）。
    色はテーマの status だけを使う。遅れている行は強調の面に「遅延」の色の線を添える。"""
    s, a = style.body(sd)
    c = style.c
    z = _z(style)
    items = _rename(as_items(sd.get("items")), "name")
    check_count(items, 1, 8, "進捗の項目")
    if not items:
        return
    n = len(items)
    late_style = style.status_style("遅延")
    todo_style = style.status_style("未着手")
    for it in items:
        st = K.as_text(it.get("status", "") or "未着手")
        if style.status_style(st) == todo_style and st not in ("未着手", "予定", "todo", "保留"):
            K.WARN.add(f"status「{st}」は 完了・進行中・遅延・未着手 のどれかにしてください"
                       "（未着手として数えます）")

    # ---- 縦の配分: 件数のまとめ（2件以上）→ 列見出し → 行。行が詰まるときは、まとめを省く
    head_h = _head_h(z)
    sum_size = z["body_small"]
    ch = _pill_h(sum_size)
    chips_h = ch + 0.22
    chips = n >= 2
    rh_max = 0.9 if n <= 2 else 0.74

    def row_h(with_chips):
        return min(rh_max, (a.h - (chips_h if with_chips else 0.0) - head_h - 0.04) / n)

    rh = row_h(chips)
    if chips and rh < 0.44:
        chips = False
        rh = row_h(False)
    compact = rh < 0.44
    # 上限を超える数のときも、文字を小さくして枠に収める（0.32 インチまで）
    rh = max(rh, 0.32)
    total = (chips_h if chips else 0.0) + head_h + 0.04 + rh * n
    _fits(a.y + total, a, "進捗の一覧")
    y = a.y

    # ---- 上: 状況ごとの件数（ラベルは正式な名前。別名・不明な状況もまとめて数える）
    if chips:
        canon = {}
        for w in reversed(_STATUS_ORDER):
            canon[style.status_style(w)] = w
        counts = {}
        for it in items:
            key = style.status_style(K.as_text(it.get("status", "") or "未着手"))
            counts[key] = counts.get(key, 0) + 1
        order = [style.status_style(w) for w in _STATUS_ORDER]
        x = a.x
        txt = f"全 {n} 件"
        w = _chip_w(txt, sum_size)
        style.pill(s, x, y, w, ch, txt, tone="soft", size=sum_size)
        x += w + 0.14
        for k in sorted(counts, key=lambda k: order.index(k) if k in order else 9):
            txt = f"{canon.get(k, '未着手')} {counts[k]}件"
            w = _chip_w(txt, sum_size)
            _chip(style, s, x, y, w, ch, txt, k[0], k[1], sum_size)
            x += w + 0.12
        y += chips_h

    # ---- 列の配置（状況のラベルを左端に置き、項目名と離れないようにする）
    has_owner = any(it.get("owner") not in (None, "") for it in items)
    has_due = any(it.get("due") not in (None, "") for it in items)
    has_note = any(it.get("note") for it in items)
    has_prog = any(it.get("progress") not in (None, "") for it in items)
    pad = 0.14
    sw = 1.2
    pgw = 2.3 if has_prog else 0.0
    ow = 1.7 if has_owner else 0.0
    dw = 0.95 if has_due else 0.0
    gaps = 0.2 * sum(1 for v in (pgw, ow, dw) if v) + 0.2
    flex = a.w - 2 * pad - sw - pgw - ow - dw - gaps - (0.2 if has_note else 0.0)
    # 項目名は本文（多いときは小さめの本文）、状況・進捗は小さめの本文、担当・期限・備考は注釈
    size = z["body"] if n <= 5 else z["body_small"]
    small = z["body_small"]
    cap = z["caption"]
    if compact:
        size = z["body_small"] if rh >= 0.38 else cap
        small = cap
    need = max(K.text_em(K.plain(K.as_text(it.get("name", ""))), True)
               for it in items) * size / 72 / SAFE + 0.1
    if has_note:
        # 備考があるときは半分ほど。項目名が1行に収まるなら、備考の幅を少し削ってでも1行にする
        name_w = min(max(need, flex * 0.48), flex * 0.62)
    else:
        # 備考がないときは、項目名を要る幅（全体の4割ほど）にとどめ、余りを進捗の棒・担当に回す
        # （どちらもなければ項目名へ。状況のラベルは左端なので、項目名が広くても離れない）
        name_w = min(flex, max(need + 0.2, 2.6, flex * 0.4))
        extra = flex - name_w
        if has_prog:
            pgw += extra
            extra = 0.0
        if has_owner:
            add = min(extra, 0.8)
            ow += add
            extra -= add
        name_w += extra
    note_w = flex - name_w if has_note else 0.0
    cols = []
    x = a.x + pad
    for key, w, title, align in (("status", sw, "状況", "c"), ("name", name_w, "項目", "l"),
                                 ("progress", pgw, "進捗", "l"), ("owner", ow, "担当", "l"),
                                 ("due", dw, "期限", "c"), ("note", note_w, "備考", "l")):
        if w > 0:
            cols.append((key, x, w, title, align))
            x += w + 0.2
    for key, cx, w, title, align in cols:
        _head(style, s, cx, y, w, head_h, title, align=align)
    y += head_h + 0.04
    K.add_line(s.shapes, a.x, y, a.right, y, _faint(c), 1.0)
    track = _grid(c)
    chip_h = min(_pill_h(small), rh - 0.08)
    page_bg = c.get("bg", WHITE)

    for i, it in enumerate(items):
        st = K.as_text(it.get("status", "") or "未着手")
        fill, fg = style.status_style(st)
        late = (fill, fg) == late_style
        bg = page_bg
        if late:
            # 遅れている行は、強調の面に「遅延」の色の線を添えて目立たせる
            bg = _attention(style, s, Area(a.x, y + 0.03, a.w, rh - 0.06), fill)
        cy = y + rh / 2
        for key, cx, w, _, align in cols:
            if key == "name":
                K.max_chars_warn(it.get("name", ""), 24, "進捗の項目名")
                style.text(s, cx, y, w, rh, K.as_text(it.get("name", "")), size, c["heading"],
                           bold=True, anchor="m", fit_min=10, label="進捗の項目名",
                           line_spacing=1.0)
            elif key == "status":
                _chip(style, s, cx + 0.06, cy - chip_h / 2, w - 0.12, chip_h, st, fill, fg,
                      small)
            elif key == "progress":
                pv = it.get("progress")
                if pv in (None, ""):
                    _blank(style, s, cx, y, w, rh, small, align="l")
                    continue
                p = _num(str(pv).strip().rstrip("%％"))
                if p is None:
                    K.WARN.add(f"progress「{pv}」は 0〜100 の数にしてください")
                    _blank(style, s, cx, y, w, rh, small, align="l")
                    continue
                if not 0 <= p <= 100:
                    K.WARN.add(f"progress {pv} は 0〜100 にしてください")
                    p = max(0.0, min(100.0, p))
                bw = w - 0.7
                bh = 0.16
                K.add_box(s.shapes, cx, cy - bh / 2, bw, bh, fill=track, geom="pill")
                if p > 0:
                    K.add_box(s.shapes, cx, cy - bh / 2, max(bh, bw * p / 100), bh,
                              fill=fill if _dark(fill) else c["primary"], geom="pill")
                style.text(s, cx + bw + 0.08, y, 0.62, rh, f"{p:.0f}%", small,
                           c["heading"], bold=True, anchor="m", latin=style.font("number_latin"))
            elif key == "owner":
                if it.get("owner") not in (None, ""):
                    style.text(s, cx, y, w, rh, K.as_text(it["owner"]), cap, c["text"],
                               anchor="m", fit_min=9, label="担当", line_spacing=1.0)
                else:
                    _blank(style, s, cx, y, w, rh, cap, align="l")
            elif key == "due":
                if it.get("due") not in (None, ""):
                    style.text(s, cx, y, w, rh, K.as_text(it["due"]), cap, c["text"], align="c",
                               anchor="m", fit_min=9, label="期限",
                               latin=style.font("number_latin"))
                else:
                    _blank(style, s, cx, y, w, rh, cap)
            elif key == "note" and it.get("note"):
                K.max_chars_warn(it["note"], 30, "進捗の備考")
                style.text(s, cx, y, w, rh, K.as_text(it["note"]), cap,
                           _ink(c, fill, bg, c["heading"]) if late else c["text_muted"],
                           bold=late, anchor="m", fit_min=9, label="進捗の備考",
                           line_spacing=1.0)
        y += rh
        # 強調した行（遅延）の上下には区切り線を引かない（枠の線と二重にならないように）
        nxt_late = i < n - 1 and style.status_style(items[i + 1].get("status")) == late_style
        if i < n - 1 and not late and not nxt_late:
            style.row_rule(s, a.x, a.right, y)
