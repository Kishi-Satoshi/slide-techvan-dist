#!/usr/bin/env python3
"""内容ファイル（deck.json）から PowerPoint を作る。

使い方:
    python build_pptx.py deck.json -o 出力.pptx

deck.json の書き方は references/slide-types.md を参照。
生成後、文字数の超過やはみ出しの可能性があれば「確認が必要な点」として表示する。

終了コード: 0 = 作成できた（確認が必要な点があっても 0）、1 = 作成できなかった
"""
import argparse
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import slidekit as K  # noqa: E402

K.console_utf8()

try:
    from pptx import Presentation  # noqa: E402
except ImportError:
    print("python-pptx が入っていません。次のコマンドで入れてください:\n"
          "  python -m pip install python-pptx")
    sys.exit(1)

from styles_impl import make_style  # noqa: E402
import renderers  # noqa: E402

TEMPLATE = HERE.parent / "assets" / "template.pptx"

APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    "<Application>Microsoft Office PowerPoint</Application>"
    "<PresentationFormat>Widescreen</PresentationFormat><Slides>{n}</Slides></Properties>"
)


class DeckError(Exception):
    """deck.json の内容の誤り（日本語のメッセージを持つ）。"""


def remove_all_slides(prs):
    """テンプレートの説明用スライドを除く（マスター・レイアウトはそのまま）。"""
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


def set_properties(prs, deck):
    cp = prs.core_properties
    cp.title = K.plain(K.as_text(deck.get("title")))
    cp.author = K.as_text(deck.get("author"))
    cp.last_modified_by = K.as_text(deck.get("author"))
    cp.subject = K.plain(K.as_text(deck.get("subtitle")))
    cp.keywords = ""
    cp.comments = ""
    cp.category = ""
    cp.revision = 1


def scrub_package(prs):
    """テンプレート由来のファイル情報を消す。

    サムネイル（テンプレートの表紙の画像）、テンプレートのスライド名の一覧（app.xml）、
    社内の文書管理の情報（custom.xml・customXml）が、作った資料に残らないようにする。
    """
    pkg = prs.part.package
    for rId, rel in list(pkg._rels.items()):
        kind = rel.reltype.rsplit("/", 1)[-1]
        if kind in ("thumbnail", "custom-properties"):
            pkg.drop_rel(rId)
        elif kind == "extended-properties":
            rel.target_part._blob = APP_XML.format(n=len(prs.slides)).encode("utf-8")
    for rId, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith("/customXml"):
            prs.part.drop_rel(rId)


def check_deck(deck):
    if not isinstance(deck, dict):
        raise DeckError("deck.json の一番外側は { … } （オブジェクト）にしてください")
    slides = deck.get("slides")
    if not isinstance(slides, list) or not slides:
        raise DeckError("deck.json に \"slides\": [ … ] （スライドのリスト）がありません")
    for i, sd in enumerate(slides, start=1):
        if not isinstance(sd, dict):
            raise DeckError(f"{i}枚目のスライドが {{ … }} （オブジェクト）になっていません")
        if not sd.get("type"):
            raise DeckError(f"{i}枚目のスライドに \"type\" がありません")


def build(deck, out_path, template=TEMPLATE):
    check_deck(deck)
    prs = Presentation(str(template))
    remove_all_slides(prs)
    set_properties(prs, deck)
    style = make_style(prs, deck)
    for i, sd in enumerate(deck["slides"], start=1):
        t = sd.get("type")
        K.WARN.where = f"{i}枚目（{t}）: "
        fn = renderers.REGISTRY.get(t)
        if fn is None:
            K.WARN.add(f"未対応の型です: {t}。references/slide-types.md の型を使ってください")
            continue
        fn(style, sd)
    K.WARN.where = ""
    style.finish()
    scrub_package(prs)
    prs.save(str(out_path))
    return K.WARN.items


def default_out(deck_path):
    """「資料.deck.json」→「資料.pptx」、「資料.json」→「資料.pptx」。"""
    name = deck_path.name
    for suffix in (".deck.json", ".json"):
        if name.endswith(suffix):
            return deck_path.with_name(name[: -len(suffix)] + ".pptx")
    return deck_path.with_suffix(".pptx")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck", help="内容ファイル deck.json")
    ap.add_argument("-o", "--out", help="出力する .pptx（既定: 資料.deck.json → 資料.pptx）")
    ap.add_argument("--template", default=str(TEMPLATE), help=argparse.SUPPRESS)
    args = ap.parse_args()

    deck_path = Path(args.deck)
    try:
        # utf-8-sig: メモ帳などで保存した BOM 付きの UTF-8 も読める
        with open(deck_path, encoding="utf-8-sig") as f:
            deck = json.load(f)
    except FileNotFoundError:
        print(f"内容ファイルが見つかりません: {deck_path}")
        return 1
    except UnicodeDecodeError:
        print(f"内容ファイルを UTF-8 として読めません: {deck_path}（UTF-8 で保存し直してください）")
        return 1
    except json.JSONDecodeError as e:
        print(f"内容ファイルの JSON の書き方に誤りがあります（{e.lineno} 行目 {e.colno} 文字目付近）: {e.msg}")
        return 1

    out = Path(args.out) if args.out else default_out(deck_path)
    try:
        out.parent.mkdir(parents=True, exist_ok=True)
        warnings = build(deck, out, args.template)
    except DeckError as e:
        print(f"作成できませんでした: {e}")
        return 1
    except PermissionError:
        print(f"保存できませんでした: {out}\n"
              "同じ名前のファイルを PowerPoint で開いていれば閉じてから、もう一度実行してください。")
        return 1
    except SystemExit:
        raise
    print(f"作成しました: {out}")
    if warnings:
        print(f"確認が必要な点（{len(warnings)} 件）:")
        for w in warnings:
            print("  - " + w)
    else:
        print("確認が必要な点はありません")
    print_text_notes(deck)
    return 0


def print_text_notes(deck):
    """文面の確認（check_text.py）。作成の成否には関係させない。"""
    try:
        import check_text
        notes = check_text.check(deck)
        prov = check_text.provisional(deck)
    except Exception as e:  # noqa: BLE001 - 文面の検査が失敗しても作成は成功のまま
        print(f"文面の確認ができませんでした（{type(e).__name__}）")
        return
    if notes:
        print(f"文面の確認（{len(notes)} 件。references/ai-smell-ja.md の言い換えの例を参考に）:")
        for n in notes:
            print("  - " + n)
    else:
        print("文面の確認: 指摘はありません")
    if prov:
        print(f"仮のまま残っている箇所（{len(prov)} 件。渡すときに一覧を添えて確かめてもらう）:")
        for p in prov:
            print("  - " + p)


if __name__ == "__main__":
    sys.exit(main())
