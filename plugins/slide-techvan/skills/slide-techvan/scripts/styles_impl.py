"""デザインの描き方（techvan）。

型（renderers.py）は「どこに何を置くか」だけを決め、
色・角丸・ラベルの形などの見た目はここで切り替える。
"""
from dataclasses import dataclass

from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

import slidekit as K


@dataclass
class Area:
    x: float
    y: float
    w: float
    h: float

    @property
    def bottom(self):
        return self.y + self.h

    @property
    def right(self):
        return self.x + self.w


def make_style(prs, deck):
    style = deck.get("style", "techvan")
    if style == "techvan":
        return TechvanStyle(prs, deck, K.load_theme("techvan"))
    raise SystemExit(f"デザイン「{style}」はありません（使えるのは techvan だけです。"
                     "\"style\" は省くか \"techvan\" にしてください）")


class BaseStyle:
    def __init__(self, prs, deck, theme):
        self.prs = prs
        self.deck = deck
        self.t = theme
        self.c = theme["colors"]
        self.f = theme["fonts"]
        self.m = theme["metrics"]
        self.layouts = self._find_layouts(theme["layouts"])
        self.current_section = {}
        self.audience = deck.get("audience", "internal")

    # ------------------------------------------------------------ 共通

    def _find_layouts(self, names):
        found = {}
        all_layouts = {}
        for master in self.prs.slide_masters:
            for layout in master.slide_layouts:
                all_layouts.setdefault(layout.name, layout)
        for kind, name in names.items():
            if name not in all_layouts:
                raise SystemExit(
                    f"テンプレートにレイアウト「{name}」がありません。"
                    "scripts/build_template.py で template.pptx を作り直してください")
            found[kind] = all_layouts[name]
        return found

    def new_slide(self, kind):
        return self.prs.slides.add_slide(self.layouts[kind])

    def cleanup(self, slide):
        """使わなかったプレースホルダーを消す（編集画面に「テキストを入力」が残らないように）。"""
        for ph in list(slide.placeholders):
            if not ph.has_text_frame or not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)

    def finish(self):
        pass

    def font(self, key="ja"):
        return self.f.get(key, K.FONT_JA)

    def shadow(self, shp):
        """影（デザインによっては付けない）。"""
        return shp

    # ------------------------------------------------------------ 型で共通に使う決まり
    #   文字サイズ: size("item_head"|"body"|"body_large"|"body_small"|"caption"|"quote", 既定)
    #   段組みの間隔: gutter()
    #   本体の置き方: 本体は style.body() が返す Area の上端（a.y）から上詰めで置く（上下中央に寄せない）
    #   番号の丸: marker(slide, x, y, text, size="m"|"l"|"s", solid=False) … 大きさは 3 段階だけ

    MARKER = {"s": (0.36, 11), "m": (0.46, 14), "l": (0.6, 18)}

    def gutter(self):
        return 0.28

    def marker(self, slide, x, y, text, size="m", solid=False, muted=False):
        """番号や Q・A などの丸。左上 (x, y) に置き、直径を返す。
        solid=True は濃い丸（目次の番号・Q など）、muted=True は控えめな丸（選ばれていない項目など）。"""
        d, pt = self.MARKER[size]
        c = self.c
        if muted:
            fill, line, color = c["bg_alt"], c.get("line_soft"), c["text_muted"]
        elif solid:
            fill, line, color = c["primary"], None, c["on_primary"]
        else:
            fill, line, color = self._marker_colors()
        b = K.add_box(slide.shapes, x, y, d, d, fill=fill, line=line, line_w=0.75, geom="oval")
        self.text(slide, 0, 0, 0, 0, K.as_text(text), pt, color, bold=True, align="c", anchor="m",
                  shape=b, margin=0, latin=self.font("number_latin"), tidy=False)
        return d

    def _marker_colors(self):
        c = self.c
        return c["bg_alt"], None, c["heading"]

    def radius(self):
        return self.m.get("radius", 0.08)

    def size(self, key, default):
        """文字サイズの段階（テーマの metrics.sizes）。なければ default。"""
        return self.m.get("sizes", {}).get(key, default)

    def text(self, slide, x, y, w, h, text, size, color=None, **kw):
        kw.setdefault("font", self.font("ja"))
        kw.setdefault("latin", self.font("latin"))
        kw.setdefault("emph_color", self.c.get("emph"))
        return K.add_text(slide.shapes, x, y, w, h, text, size, color or self.c["text"], **kw)

    def content_area(self, top):
        m = self.m
        return Area(m["margin_x"], top, K.SLIDE_W - 2 * m["margin_x"], m["content_bottom"] - top)

    def reserve_bottom(self, slide, area, sd):
        """callout（まとめの一文）と note（注記）の場所を下から確保して描く。"""
        gap = 0.16
        if sd.get("note"):
            K.max_chars_warn(sd["note"], 120, "注記")
            h = 0.34
            self.note(slide, Area(area.x, area.bottom - h, area.w, h), sd["note"])
            area = Area(area.x, area.y, area.w, area.h - h - 0.06)
        if sd.get("callout"):
            K.max_chars_warn(sd["callout"], 60, "まとめの一文（callout）")
            h = self.m.get("callout_h", 0.5)
            self.callout(slide, Area(area.x, area.bottom - h, area.w, h), sd["callout"])
            area = Area(area.x, area.y, area.w, area.h - h - gap)
        return area

    def note(self, slide, a, text):
        self.text(slide, a.x, a.y, a.w, a.h, "※ " + text, 9, self.c["text_muted"], anchor="m",
                  fit_min=8, label="注記")

    # ------------------------------------------------------------ 共通の部品（各デザインで上書き）
    #
    # 型（renderers*.py）はここにある部品だけを使って描く。見た目はデザインごとに違う。
    #   plain_card(slide, a, idx=None, highlight=False)  … 中身を入れる箱（カード）
    #   panel(slide, a)                                  … 淡い色の面（補足・まとめ）
    #   pill(slide, x, y, w, h, text, tone, size)        … ラベル（tone: primary / soft / accent0〜3）
    #   badge(slide, x, y, d, text)                      … 番号の丸
    #   sub_heading(slide, x, y, w, text, size)          … ブロックの小見出し（高さを返す）
    #   arrow(slide, x, y, w, h, direction)              … 矢印（right / down）
    #   level_style(level)                               … 高・中・低 の色 → (塗り, 文字色)
    #   status_style(status)                             … 完了・進行中・遅延・未着手 の色 → (塗り, 文字色)
    #   series_colors()                                  … グラフの系列の色

    def plain_card(self, slide, a, idx=None, highlight=False):
        c = self.c
        return K.add_box(slide.shapes, a.x, a.y, a.w, a.h,
                         fill=c["bg_alt"] if highlight else c["card_bg"],
                         line=c["card_border"], line_w=1.0, geom="round", radius=0.08)

    def panel(self, slide, a):
        return K.add_box(slide.shapes, a.x, a.y, a.w, a.h, fill=self.c["bg_alt"], geom="round",
                         radius=0.08)

    def pill(self, slide, x, y, w, h, text, tone="primary", size=11, align="c"):
        c = self.c
        if tone == "soft":
            fill, line, color = c["bg_alt"], c.get("line_soft") or c["card_border"], c["heading"]
        elif tone.startswith("accent"):
            accents = c["series"]
            fill, line, color = accents[int(tone[6:] or 0) % len(accents)], None, "FFFFFF"
        else:
            fill, line, color = c["primary"], None, c["on_primary"]
        b = K.add_box(slide.shapes, x, y, w, h, fill=fill, line=line, line_w=0.75, geom="pill")
        self.text(slide, 0, 0, 0, 0, text, size, color, bold=True, align=align, anchor="m",
                  shape=b, margin=0.06, fit_min=max(7.5, size - 3), label="ラベル")
        return b

    def sub_heading(self, slide, x, y, w, text, size=14):
        h = max(0.34, size * K.LINE_HEIGHT / 72 + 0.04)
        self.text(slide, x, y, w, h, text, size, self.c["heading"], bold=True, anchor="m",
                  fit_min=max(10, size - 3), label="小見出し")
        return h

    def arrow(self, slide, x, y, w, h, direction="right", color=None):
        geom = "chevron" if direction == "right" else "triangle"
        shp = K.add_box(slide.shapes, x, y, w, h, fill=color or self.c["primary"], geom=geom,
                        rotation=180 if direction == "down" else 0)
        return shp

    def level_style(self, level):
        lv = str(level).strip()
        levels = self.c["levels"]
        if lv in ("高", "大", "high", "最優先", "致命", "最大"):
            return tuple(levels["high"])
        if lv in ("中", "mid", "medium"):
            return tuple(levels["mid"])
        return tuple(levels["low"])

    def status_style(self, status):
        c = self.c
        st = str(status)
        table = {
            "done": ("完了", "済", "done", "ok", "達成"),
            "doing": ("進行中", "対応中", "doing", "on track", "順調"),
            "late": ("遅延", "遅れ", "要対応", "late", "ng", "未達", "注意"),
            "todo": ("未着手", "予定", "todo", "保留"),
        }
        for key, words in table.items():
            if st in words:
                return tuple(c["status"][key])
        return tuple(c["status"]["todo"])

    def series_colors(self):
        return list(self.c["series"])

    # ------------------------------------------------------------ 型から呼ばれる部品（既定）

    def kpi_value(self, slide, a, value, unit, size):
        self.number_runs(slide, a, value, unit, size, self.c["number"])

    def kpi_note(self, slide, a, text):
        self.text(slide, a.x + 0.2, a.y, a.w - 0.4, a.h, text, 12, self.c["text_muted"],
                  align="c", anchor="m", fit_min=9, label="数字の補足")

    def row_rule(self, slide, x1, x2, y):
        K.add_line(slide.shapes, x1, y, x2, y, self.c["rule"], 0.75)

    # ------------------------------------------------------------ 数字の強調

    def number_runs(self, slide, a, value, unit, size, color, align="c", unit_ratio=0.42,
                    anchor="m"):
        """大きな数字＋小さな単位（例: 1,824.2 億円）。枠に収まるよう縮める。"""
        value = str(value)
        unit = str(unit or "")
        # 数字と単位を合わせた幅で大きさを決める
        while size > 14:
            w = (K.text_em(value, True) * size + K.text_em(unit, True) * size * unit_ratio) / 72.0
            if w <= a.w - 0.1 and size * K.LINE_HEIGHT / 72.0 <= a.h + 0.05:
                break
            size -= 1
        box = slide.shapes.add_textbox(Inches(a.x), Inches(a.y), Inches(a.w), Inches(a.h))
        tf = box.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = K.ANCHOR[anchor]
        p = tf.paragraphs[0]
        K.set_paragraph(p, align, 1.0)
        r = p.add_run()
        r.text = value
        K.style_run(r, size=size, bold=True, color=color, font=self.font("ja"),
                    latin=self.font("number_latin"))
        if unit:
            r2 = p.add_run()
            r2.text = unit
            K.style_run(r2, size=max(10, round(size * unit_ratio)), bold=True, color=color,
                        font=self.font("ja"), latin=self.font("number_latin"))
        return box


# ====================================================================== techvan

class TechvanStyle(BaseStyle):
    """テクバン: 紺と白。英字アイブロウ → 見出し → リード文 → 本体。

    寸法は社内の分析資料（16:9）の実測値に合わせている。
    """

    def shadow(self, shp):
        K.set_soft_shadow(shp, alpha=35, blur=6 / 72, dist=2 / 72, color=self.c["shadow"])
        return shp

    def _marker_colors(self):
        c = self.c
        return c["bg_alt"], c["line_soft"], c["heading"]

    def row_rule(self, slide, x1, x2, y):
        # 淡い青の面の上でも見える線の色
        K.add_line(slide.shapes, x1, y, x2, y, self.c["line_soft"], 0.75)

    def white_card(self, slide, x, y, w, h, fill=None, line=None):
        shp = K.add_box(slide.shapes, x, y, w, h, fill=fill or self.c["card_bg"],
                        line=line or self.c["card_border"], line_w=1.0, geom="round", radius=0.06)
        self.shadow(shp)
        return shp

    # ------------------------------------------------------------ 機密のマーク

    CONFIDENTIAL_LEVELS = {
        "internal": "internal", "社外秘": "internal", "confidential": "internal",
        "strict": "strict", "厳秘": "strict", "secret": "strict",
        "none": None, "なし": None, "": None,
    }

    def confidential_level(self, sd, default_key=True):
        """そのページの機密の区分（internal / strict / None）。ページの指定 → 資料全体の指定の順。"""
        raw = sd.get("confidential")
        if raw is None and default_key:
            raw = self.deck.get("confidential")
        if raw is None or raw is False:
            return None
        if raw is True:
            return "internal"
        key = str(raw).strip().lower() if str(raw).isascii() else str(raw).strip()
        if key not in self.CONFIDENTIAL_LEVELS:
            K.WARN.add(f"confidential は internal（社外秘）/ strict（厳秘）/ none のどれかにしてください（{raw}）")
            return None
        return self.CONFIDENTIAL_LEVELS[key]

    def confidential_mark(self, slide, level, x, y, w=1.95, h=0.38):
        """CONFIDENTIAL のマーク。internal = 赤枠の白地、strict = 赤地に白文字。"""
        red = self.c["confidential"]
        if level == "strict":
            box = K.add_box(slide.shapes, x, y, w, h, fill=red, name="CONFIDENTIAL")
            color = "FFFFFF"
        else:
            box = K.add_box(slide.shapes, x, y, w, h, fill="FFFFFF", line=red, line_w=2.0,
                            name="CONFIDENTIAL")
            color = red
        self.text(slide, 0, 0, 0, 0, "CONFIDENTIAL", 15, color, bold=True, align="c", anchor="m",
                  shape=box, margin=0.02, font="Arial", latin="Arial", spacing=0.5, tidy=False)
        return box

    # ------------------------------------------------------------ 表紙・章扉・最終ページ

    COVER_MARK_XY = (0.9, 0.75)   # 表紙のマークの位置（カードの左上）

    def cover(self, sd):
        s = self.new_slide("cover")
        c = self.c
        cx, w = 0.9, 9.8
        y = 2.15
        if sd.get("badge"):
            badge = " ".join(list(sd["badge"])) if len(sd["badge"]) <= 4 else sd["badge"]
            bw = max(2.1, K.text_em(badge, True) * 15 / 72 + 0.7)
            b = K.add_box(s.shapes, cx + (w - bw) / 2, y, bw, 0.5, fill=c["primary"], geom="pill")
            self.text(s, 0, 0, 0, 0, badge, 15, c["on_primary"], bold=True, align="c",
                      anchor="m", shape=b, spacing=4, margin=0.04)
        y = 2.85
        title = sd.get("title", self.deck.get("title", ""))
        K.max_chars_warn(title, 26, "表紙のタイトル")
        ph = s.shapes.title
        lines = 2 if "\n" in title or K.count_lines(K.plain(title), w - 0.2, 34, True) > 1 else 1
        th = 0.8 if lines == 1 else 1.4
        K.place(ph, cx, y, w, th)
        K.fill_placeholder(ph, title, 34, c["primary"], bold=True, align="c", anchor="m",
                           font=self.font("ja"), latin=self.font("latin"), fit_min=24,
                           emph_color=c["emph"], label="表紙のタイトル", marker=c.get("marker"))
        y += th + 0.1
        if sd.get("subtitle"):
            self.text(s, cx, y, w, 0.4, "― " + sd["subtitle"] + " ―", 14, c["text"],
                      align="c", anchor="m", fit_min=11, label="サブタイトル", marker=c.get("marker"))
            y += 0.65
        K.add_line(s.shapes, cx + w / 2 - 1.65, y, cx + w / 2 + 1.65, y, c["line_soft"], 1.0)
        y += 0.2
        if sd.get("client"):
            self.text(s, cx, y, w, 0.4, sd["client"], 14, c["text"], align="c", anchor="m",
                      fit_min=11)
            y += 0.65
        meta = [sd.get("date", self.deck.get("date", "")), sd.get("author", self.deck.get("author", ""))]
        meta_txt = "\n".join(t for t in meta if t)
        if meta_txt:
            self.text(s, cx, y, w, 0.75, meta_txt, 12.5, c["heading"], align="c", anchor="t",
                      line_spacing=1.4)
        # 表紙は、ページか資料全体に機密の指定があるときだけマークを付ける
        level = self.confidential_level(sd)
        if level:
            self.confidential_mark(s, level, *self.COVER_MARK_XY)
        self.cleanup(s)
        return s

    def section(self, sd):
        dark = sd.get("variant") == "dark"
        s = self.new_slide("section_dark" if dark else "section")
        c = self.c
        self.current_section = {"title": sd.get("title", ""), "en": sd.get("en", ""),
                                "ja": sd.get("title", "")}
        title = sd.get("title", "")
        if sd.get("number"):
            title = f"{sd['number']}　{title}"
        K.max_chars_warn(title, 20, "章のタイトル")
        ph = s.shapes.title
        if dark:
            K.fill_placeholder(ph, title, 32, "FFFFFF", bold=True, anchor="m", font=self.font("ja"),
                               fit_min=20, label="章のタイトル")
            if sd.get("subtitle"):
                self.text(s, 1.57, 3.85, 8.6, 0.45, sd["subtitle"], 16, "FFFFFF", fit_min=11)
        else:
            cx, w = 1.2, 7.5
            K.place(ph, cx, 2.9, w, 0.9)
            K.fill_placeholder(ph, title, 40, c["primary"], bold=True, align="c", anchor="m",
                               font=self.font("ja"), fit_min=24, label="章のタイトル")
            if sd.get("subtitle"):
                self.text(s, cx, 3.85, w, 0.45, sd["subtitle"], 16, c["text"], align="c",
                          anchor="m", fit_min=11, label="章のサブタイトル")
            if sd.get("en"):
                self.text(s, cx, 4.4, w, 0.4, "-  " + " ".join(sd["en"]) + "  -", 13,
                          c["text_faint"], align="c", anchor="m", latin=self.font("eyebrow_latin"),
                          spacing=2)
        self.cleanup(s)
        return s

    def closing(self, sd):
        """テンプレートの最終ページ（Thank you と会社情報）をそのまま使う。
        message・contact は使わない（写真と文字に重なるため）。"""
        s = self.new_slide("closing")
        if sd.get("message") or sd.get("contact"):
            K.WARN.add("techvan の最終ページはテンプレートのまま（Thank you と会社情報）なので、"
                       "message・contact は表示されません")
        self.cleanup(s)
        return s

    # ------------------------------------------------------------ 本文ページ

    def body(self, sd):
        """英字ラベル（y0.28）→ 見出し（y0.56）→ リード文（y1.14）→ 本体（y1.72〜）。"""
        s = self.new_slide("body")
        c = self.c
        m = self.m
        x = m["margin_x"]
        w = K.SLIDE_W - 2 * x
        en = sd.get("eyebrow") or self.current_section.get("en", "")
        ja = sd.get("eyebrow_ja") or sd.get("label") or self.current_section.get("ja", "")
        parts = []
        if en:
            parts.append(en.upper())
        if ja:
            parts.append("-  " + K.spaced_ja(ja) + "  -")
        level = self.confidential_level(sd)
        mark_w = 1.95
        if level:
            self.confidential_mark(s, level, K.SLIDE_W - x - mark_w, 0.17, mark_w)
        if parts:
            self.text(s, x, 0.28, w - (mark_w + 0.2 if level else 0), 0.28, "  ".join(parts),
                      10.5, c["primary"], bold=True,
                      latin=self.font("eyebrow_latin"), spacing=3, anchor="m", tidy=False)
        title = sd.get("title", "")
        K.max_chars_warn(title, 36, "見出し")
        K.fill_placeholder(s.shapes.title, title, m["title_size"], c["heading"], bold=True,
                           anchor="m", font=self.font("ja"), latin=self.font("latin"),
                           emph_color=c["emph"], fit_min=16, label="見出し", marker=c.get("marker"))
        top = 1.42
        if sd.get("lead"):
            K.max_chars_warn(sd["lead"], 80, "リード文")
            self.text(s, x, 1.14, w, 0.3, sd["lead"], 12, c["text"], anchor="m", fit_min=10,
                      label="リード文", marker=c.get("marker"))
            top = 1.72
        area = self.content_area(top)
        area = self.reserve_bottom(s, area, sd)
        self.cleanup(s)
        return s, area

    def callout(self, slide, a, text):
        box = K.add_box(slide.shapes, a.x, a.y, a.w, a.h, fill=self.c["primary"],
                        line=self.c["primary"], line_w=1.0, geom="round", radius=0.06)
        self.text(slide, 0, 0, 0, 0, text, 13, "FFFFFF", bold=True, align="c", anchor="m",
                  shape=box, margin=0.1, fit_min=10.5, label="まとめの一文",
                  emph_color="FFFFFF")

    def note(self, slide, a, text):
        """注記は飾りのない小さな文字（淡い箱はまとめの一文用にとっておく）。"""
        self.text(slide, a.x, a.y, a.w, a.h, "※ " + text, self.size("note", 9.5),
                  self.c["text_muted"], anchor="m", fit_min=8, label="注記")

    # ------------------------------------------------------------ 部品

    def badge(self, slide, x, y, d, text):
        """淡い青の丸に紺の数字（アイコンの丸と同じ形）。"""
        b = K.add_box(slide.shapes, x, y, d, d, fill=self.c["bg_alt"], line=self.c["line_soft"],
                      line_w=0.75, geom="oval")
        self.text(slide, 0, 0, 0, 0, str(text), max(9, d * 30), self.c["heading"], bold=True,
                  align="c", anchor="m", shape=b, margin=0, latin=self.font("number_latin"))
        return b

    def bullet_marker(self, slide, x, y, size, idx, item):
        return self.badge(slide, x, y, size, item.get("mark", idx + 1))

    def card_header_h(self, item, w):
        return 0.8

    @staticmethod
    def _card_title(item):
        """番号の丸と同じ番号を見出しで繰り返さない（「理由1」などの番号付きラベルは省く）。"""
        label = item.get("label", "")
        heading = item.get("heading", "")
        if label and not label[-1].isdigit():
            return f"{label}　{heading}"
        return heading

    def card(self, slide, a, item, idx, highlight=False):
        c = self.c
        self.plain_card(slide, a, highlight=highlight)
        pad = 0.24
        y = a.y + 0.2
        d = self.marker(slide, a.x + pad, y, item.get("mark", idx + 1), "m")
        hx = a.x + pad + d + 0.14
        hw = a.right - hx - 0.18
        self.text(slide, hx, y, hw, d, self._card_title(item), self.size("card_title", 14),
                  c["heading"], bold=True, anchor="m", fit_min=10.5, label="カードの見出し",
                  line_spacing=0.95)
        y = a.y + 0.84
        return Area(a.x + pad, y, a.w - 2 * pad, a.bottom - y - 0.18)

    def plain_card(self, slide, a, idx=None, highlight=False):
        c = self.c
        if highlight:
            return self.white_card(slide, a.x, a.y, a.w, a.h, fill=c["tint"], line=c["line_soft"])
        return self.white_card(slide, a.x, a.y, a.w, a.h)

    def panel(self, slide, a):
        return K.add_box(slide.shapes, a.x, a.y, a.w, a.h, fill=self.c["bg_alt"],
                         line=self.c["card_border"], line_w=1.0, geom="round", radius=0.06)

    def sub_heading(self, slide, x, y, w, text, size=13):
        h = max(0.32, size * K.LINE_HEIGHT / 72)
        self.text(slide, x, y, w, h, text, size, self.c["heading"], bold=True, anchor="m",
                  fit_min=max(10, size - 3), label="小見出し")
        return h

    def kpi_tile(self, slide, a, item, idx):
        c = self.c
        self.white_card(slide, a.x, a.y, a.w, a.h)
        hh = 0.44
        band = K.add_box(slide.shapes, a.x, a.y, a.w, hh, fill=c["primary"], geom="round_top",
                         radius=0.06)
        self.text(slide, 0, 0, 0, 0, item.get("label", ""), 12, "FFFFFF", bold=True, align="c",
                  anchor="m", shape=band, fit_min=9, label="数字の項目名")
        return Area(a.x, a.y + hh, a.w, a.h - hh)

    def kpi_value(self, slide, a, value, unit, size):
        self.number_runs(slide, a, value, unit, size, self.c["number"], unit_ratio=0.36)

    def table_colors(self):
        c = self.c
        return {"head_bg": c["primary"], "head_text": "FFFFFF", "row_bg": c["bg_alt"],
                "row_alt": "FFFFFF", "line": c["card_border"], "text": c["text"],
                "first_col": c["heading"]}
