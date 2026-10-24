"""構成の型（compare / options / matrix / process / tree）。

renderers.py が自動で読み込む。配置だけをここで決め、色や角丸などの見た目は
styles_impl.py の共通の部品（plain_card・pill・row_rule など）と色の役割（style.c）に任せる。
型で共通の決まり（styles_impl.py の BaseStyle「型で共通に使う決まり」）に従う:
本体は style.body() の Area の上端から上詰めで置き、余りは下に残す。
文字の大きさは style.size() の段階（item_head / body / body_large / body_small / caption）だけを使い、
収まらないときだけ縮める（同じ種類の文字は1枚の中で同じ大きさにそろえる）。
横に並べる列は a.x から a.right までを style.gutter() の間隔で分ける。
"""
import math

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Pt

import slidekit as K
from styles_impl import Area
from renderers import register, check_count


# ====================================================================== 小さな部品（この型だけで使う）

SAFE = 0.95   # 文字幅の見積もりに持たせる余裕（実際のフォントは見積もりより少し広いことがある）


def _s(v):
    """文字にそろえる（None は空の文字）。"""
    return "" if v is None else str(v)


def _list(v):
    """リストにそろえる（1つだけ書かれていてもリストにする）。"""
    if v is None or v == "":
        return []
    return list(v) if isinstance(v, (list, tuple)) else [v]


def _text_list(v):
    """箇条書きの文字のリスト（文字・{text: …}・数のどれで書かれていてもよい）。"""
    return [_s(p.get("text", "")) if isinstance(p, dict) else _s(p) for p in _list(v)
            if p is not None]


def _index(v, n):
    """0〜n-1 の番号なら v、そうでなければ None（True / False は番号とみなさない）。"""
    return v if isinstance(v, int) and not isinstance(v, bool) and 0 <= v < n else None


def _scale(style):
    """文字の大きさの段階（テーマの metrics.sizes）。型の中では大きさの数字を直接書かない。"""
    return {k: style.size(k, d) for k, d in (("item_head", 16), ("body", 14), ("body_large", 16),
                                             ("body_small", 12), ("caption", 11))}


def _uniq(sizes):
    """大きい順に重なりなく並べる（段階が同じ値のテーマでも同じ大きさを2回試さない）。"""
    return sorted(set(sizes), reverse=True)


def _strong(c):
    """濃い塗り（見出し帯・ラベル）の色と、その上の文字色。"""
    return c["primary"], c["on_primary"]


def _lum(hex_str):
    """色の明るさ（WCAG の相対輝度）。"""
    def ch(v):
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(int(hex_str.lstrip("#")[i:i + 2], 16) / 255) for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _contrast(a, b):
    """2つの色のコントラスト比（1〜21）。"""
    la, lb = _lum(a), _lum(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def _ink(c, color, bg=None):
    """淡い地に置く色つきの文字の色。地との差が 4.5:1 に届かなければ見出しの色にする。"""
    bg = bg or c.get("card_bg", c["bg"])
    return color if _contrast(color, bg) >= 4.5 else c["heading"]


def _fill_of(shp, default):
    """図形の塗りの色（読めなければ default）。"""
    try:
        return str(shp.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return default


def _rule(style, s, x1, x2, y, bg):
    """区切りの線。地の色と近くて見えないときは、淡い文字の色の細い線にする。"""
    c = style.c
    if _contrast(c["rule"], bg) >= 1.5:
        style.row_rule(s, x1, x2, y)
    else:
        K.add_line(s.shapes, x1, y, x2, y, c.get("text_faint", c["text_muted"]), 0.75)


def _recolor(shp, fill=None, line=None, line_w=None):
    """共通の部品で作った箱の色だけを変える（形・影はそのまま）。"""
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = K.rgb(fill)
    if line:
        shp.line.color.rgb = K.rgb(line)
    if line_w:
        shp.line.width = Pt(line_w)
    return shp


def _outline(style, s, a, color, width=2.0):
    """塗りなしの枠だけを置く（強調の枠）。"""
    return K.add_box(s.shapes, a.x, a.y, a.w, a.h, line=color, line_w=width, geom="round",
                     radius=style.radius())


def _arrow_line(s, x1, y1, x2, y2, color, width=1.5):
    """終点に矢じりのある線（マトリクスの軸）。"""
    ln = K.add_line(s.shapes, x1, y1, x2, y2, color, width)
    el = ln.line._get_or_add_ln()
    etree.SubElement(el, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return ln


VLS = 0.9   # 縦に積んだ文字の行間


def _vchars(text):
    """縦に並べる文字（長音は縦棒に）。"""
    return [("｜" if ch in "ーｰ-—" else ch) for ch in K.plain(text) if not ch.isspace()]


def _vtext(style, s, x, y, w, h, text, size, color, bold=True):
    """縦書きの文字（マトリクスの縦軸の名前）。1文字ずつ改行して積む
    （縦書きの設定は LibreOffice で文字が重なるため使わない）。長いときは縮める。"""
    chars = _vchars(text)
    while size > 9 and len(chars) * size * K.LINE_HEIGHT * VLS / 72 > h - 0.16:
        size -= 0.5
    return style.text(s, x, y, w, h, "\n".join(chars), size, color, bold=bold, align="c",
                      anchor="m", line_spacing=VLS)


def _pill_h(size):
    """ラベル（pill）の高さ。pill の内側の余白をとっても文字が縮まない高さにする。"""
    return size * K.LINE_HEIGHT / 72 + 0.14


def _pill_w(text, size=12, lo=0.9, hi=4.0):
    return max(lo, min(hi, K.text_em(K.plain(text), True) * size / 72 + 0.45))


def _em(text, bold=False):
    return K.text_em(K.plain(text), bold)


def _lines_h(text, w, size, bold=False, max_lines=None, line_spacing=1.0):
    """幅 w に流したときの高さの見積もり（幅は SAFE だけ狭く見る）。"""
    n = K.count_lines(K.plain(text), w * SAFE, size, bold)
    if max_lines:
        n = min(n, max_lines)
    return n * size * K.LINE_HEIGHT * line_spacing / 72


def _balanced_w(text, w, size, bold=False):
    """折り返すときに各行の長さがそろう文字の枠の幅（1〜2 文字だけ次の行に送らない）。"""
    plain = K.plain(text)
    n = K.count_lines(plain, w * SAFE, size, bold)
    if n <= 1:
        return w
    bw = (math.ceil(K.text_em(plain, bold) / n) + 1) * size / 72 / SAFE
    return w if bw >= w or K.count_lines(plain, bw * SAFE, size, bold) > n else bw


def _points_h(points, w, size, para_space=4, line_spacing=1.05):
    """箇条書き（行頭記号つき）の高さの見積もり。"""
    if not points:
        return 0.0
    h = sum(_lines_h(p, w - 0.22, size, line_spacing=line_spacing) for p in points)
    return h + (len(points) - 1) * para_space / 72


def _fit(texts, w, h, hi, lo, bold=False, line_spacing=1.0):
    """どの文字も幅 w・高さ h に収まる、いちばん大きな文字の大きさ（そろえて使う。最小 lo）。"""
    for size in _steps(hi, lo):
        if all(_lines_h(t, w, size, bold, line_spacing=line_spacing) <= h + 0.01 for t in texts):
            return size
    return lo


def _steps(hi, lo):
    """文字の大きさの候補（hi から lo まで 0.5pt きざみ）。"""
    out, v = [], float(hi)
    while v >= lo - 1e-6:
        out.append(v)
        v -= 0.5
    return out


# ====================================================================== compare（2列の比較）

def _side(d):
    """left / right をそろえる。文字だけなら見出し、リストなら箇条書きとみなす。"""
    if d is None or d == "":
        return None
    if isinstance(d, str):
        return {"label": "", "heading": d, "body": "", "points": []}
    if isinstance(d, (list, tuple)):
        return {"label": "", "heading": "", "body": "", "points": _text_list(d)}
    if not isinstance(d, dict):
        K.WARN.add("compare の left / right は {label, heading, body, points} の形で書いてください")
        return None
    out = {k: _s(d.get(k)) for k in ("label", "heading", "body")}
    out["points"] = _text_list(d.get("points"))
    return out


@register("compare")
def compare(style, sd):
    """2列の比較（左→右の変化）。left / right: {label, heading, body, points}、
    arrow: 矢印に添える短い言葉。右側（あるべき姿）を強調する。"""
    s, a = style.body(sd)
    c = style.c
    sides = [_side(sd.get("left")), _side(sd.get("right"))]
    if not all(sides):
        K.WARN.add("compare には left と right の両方が必要です")
        sides = [x or {"label": "", "heading": "", "body": "", "points": []} for x in sides]
    arrow = _s(sd.get("arrow"))
    K.max_chars_warn(arrow, 8, "矢印の言葉")
    for side in sides:
        K.max_chars_warn(side["heading"], 30, "比較の見出し")
        K.max_chars_warn(side["body"], 60, "比較の本文")
        limit = 2 if side["body"] else 4
        if len(side["points"]) > limit:
            K.WARN.add(f"比較の箇条書きは {limit} 個までにしてください（いまは {len(side['points'])} 個"
                       + ("。本文があるときは 2 個まで）" if side["body"] else "）"))
        for p in side["points"]:
            K.max_chars_warn(p, 25, "比較の箇条書き")
    z = _scale(style)
    hs, ls = z["item_head"], z["body_small"]
    mid = max(1.3, min(2.0, _em(arrow, True) * ls / 72 + 0.35)) if arrow else 0.95
    w = (a.w - mid) / 2

    def dims(compact):
        """余白・ラベル・区切りの間。compact は詰めた組み方（余白・間を狭く、ラベルを小さく）。"""
        pad = 0.22 if compact else 0.3
        lsz = z["caption"] if compact else ls
        return {"pad": pad, "iw": w - 2 * pad, "gap": 0.08 if compact else 0.12,
                "lh": _pill_h(lsz), "ls": lsz, "lgap": 0.1 if compact else 0.14}

    def need(side, size, compact):
        d = dims(compact)
        h = 2 * d["pad"]
        if side["label"]:
            h += d["lh"] + d["lgap"]
        if side["heading"]:
            h += _lines_h(side["heading"], d["iw"], hs, True, 3) + 0.06 + 2 * d["gap"]
        if side["body"]:
            h += _lines_h(side["body"], d["iw"], size, line_spacing=1.08) + 0.12
        return h + _points_h(side["points"], d["iw"], size, 6, 1.08)

    # 本文は 2 つの側なのでゆったり（body_large）→ body → body_small。
    # それでも収まらなければ余白を詰めて 10pt まで縮める
    plans = [(v, False) for v in _uniq((z["body_large"], z["body"], z["body_small"]))]
    plans += [(v, True) for v in _steps(z["body_small"], 10)]
    for size, compact in plans:
        if max(need(x, size, compact) for x in sides) <= a.h:
            break
    else:
        K.WARN.add("比較の文字が多すぎて枠に収まりません。箇条書きを減らすか、文を短くしてください")
    d = dims(compact)
    pad, iw = d["pad"], d["iw"]
    # カードは本体の上端から。高さは中身に合わせ（低すぎると貧弱なので 2.3 インチ以上）、余りは下に残す
    box_h = min(a.h, max(max(need(x, size, compact) for x in sides) + 0.2, min(a.h, 2.3)))
    y0 = a.y
    fill, on = _strong(c)
    for k, side in enumerate(sides):
        right = k == 1
        x = a.x + k * (w + mid)
        card = style.plain_card(s, Area(x, y0, w, box_h), highlight=right)
        if right:
            _recolor(card, line=fill, line_w=2.0)
        bg = _fill_of(card, c["card_bg"])
        y = y0 + pad
        bottom = y0 + box_h - pad
        if side["label"]:
            style.pill(s, x + pad, y, _pill_w(side["label"], d["ls"], 1.1, iw), d["lh"],
                       side["label"], "primary" if right else "soft", d["ls"])
            y += d["lh"] + d["lgap"]
        if side["heading"]:
            hh = _lines_h(side["heading"], iw, hs, True, 3) + 0.06
            style.text(s, x + pad, y, iw, hh, side["heading"], hs, c["heading"], bold=True,
                       anchor="m", fit_min=max(11, hs - 4), label="比較の見出し", line_spacing=1.0)
            y += hh + d["gap"]
            _rule(style, s, x + pad, x + w - pad, y, bg)
            y += d["gap"]
        if side["body"]:
            bh = _lines_h(side["body"], iw, size, line_spacing=1.08) + 0.06
            style.text(s, x + pad, y, iw, min(bh, bottom - y), side["body"], size, c["text"],
                       anchor="t", fit_min=9, label="比較の本文", line_spacing=1.08)
            y += bh + 0.06
        if side["points"]:
            style.text(s, x + pad, y, iw, max(0.2, bottom - y + 0.04), side["points"], size,
                       c["text"], anchor="t", fit_min=9, label="比較の箇条書き", bullet="・",
                       bullet_color=fill, para_space=6, line_spacing=1.08)
    # 中央の矢印（カードの上下中央。言葉は矢印の下）
    cx = a.x + w + mid / 2
    ah = 0.86
    ay = y0 + box_h / 2 - ah / 2 - (0.22 if arrow else 0)
    style.arrow(s, cx - 0.26, ay, 0.52, ah, "right", color=fill)
    if arrow:
        style.text(s, cx - mid / 2 + 0.06, ay + ah + 0.08, mid - 0.12, 0.6, arrow, ls,
                   _ink(c, fill, c["bg"]), bold=True, align="c", anchor="t", fit_min=9,
                   label="矢印の言葉")


# ====================================================================== options（案の比較）

# 記号の色は文字の色の役割だけ（強調はおすすめ列の面とラベルに任せ、記号では目立たせない）
SYMBOLS = {"◎": "heading", "○": "heading", "△": "text_muted", "×": "text_muted"}


@register("options")
def options(style, sd):
    """案の比較表。options: 案の名前、criteria: 評価軸、scores[i][j]: 評価軸 i × 案 j、
    comments: 案ごとの一言（任意）、recommended: おすすめの案（0 から数える）。"""
    s, a = style.body(sd)
    c = style.c
    names = [_s(o) for o in _list(sd.get("options"))]
    crit = [_s(x) for x in _list(sd.get("criteria"))]
    raw = _list(sd.get("scores"))
    comments = [_s(x) for x in _list(sd.get("comments"))]
    check_count(names, 2, 4, "案（options）")
    check_count(crit, 1, 6, "評価軸（criteria）")
    if not names or not crit:
        return
    n, m = len(names), len(crit)
    if len(raw) != m or any(not isinstance(r, (list, tuple)) or len(r) != n for r in raw):
        K.WARN.add(f"scores は 評価軸 {m} 行 × 案 {n} 列 にしてください")
    scores = [list(r) if isinstance(r, (list, tuple)) else [] for r in raw]
    if comments and len(comments) != n:
        K.WARN.add(f"comments は案と同じ {n} 個にしてください（いまは {len(comments)} 個）")
    for t in names:
        K.max_chars_warn(t, {2: 16, 3: 12}.get(n, 10), "案の名前")
    for t in comments:
        K.max_chars_warn(t, 24 if n <= 3 else 16, "案の一言")
    for t in crit:
        K.max_chars_warn(t, 10, "評価軸")
    rec = sd.get("recommended")
    if rec is not None and _index(rec, n) is None:
        K.WARN.add(f"recommended は 0〜{n - 1} の数にしてください")
        rec = None

    z = _scale(style)
    gap = style.gutter()
    lw = max(1.5, min(2.8, max(_em(t, True) for t in crit) * z["body_large"] / 72 + 0.35))
    # おすすめの枠は列の外へ横 ox・上下 ov だけ出す（上下は同じだけ）。
    # ラベルは枠の上の線にまたがるので、見出しの箱に重ならないよう ov をラベルの半分より大きくする
    ox = 0.1
    pill_h = z["body_small"] * K.LINE_HEIGHT / 72 + 0.08
    ov = pill_h / 2 + 0.05
    # 列は a.right で終わる（評価軸の列＋間隔＋案の列）。右端の案がおすすめのときは、枠が a.right で終わる
    ow = (a.w - lw - gap * n - (ox if rec == n - 1 else 0)) / n
    x0 = a.x + lw + gap
    room_t, room_b = (ov + pill_h / 2, ov) if rec is not None else (0, 0)
    tw = ow - 0.3

    def header(compact):
        """見出し（案の名前＋一言）の文字の大きさと高さ。大きさは全列でそろえる。
        compact は名前を body・一言を caption にし、名前が1行に収まるなら1行にする。"""
        ns = z["body"] if compact else z["item_head"]
        cs = z["caption"] if compact else z["body_small"]
        lines = 2
        if compact and _fit(names, tw, ns * K.LINE_HEIGHT / 72, ns, 10, True) >= ns - 1.5:
            lines = 1
        name_h = max(_lines_h(t, tw, ns, True, lines) for t in names)
        ns = _fit(names, tw, name_h, ns, 10, True)
        com_h = max([_lines_h(t, tw, cs, False, 2) for t in comments if t] or [0])
        cs = min(_fit([t for t in comments if t], tw, com_h, cs, 9), ns)
        hh = max(0.5 if compact else 0.7, 0.28 + name_h + (0.06 + com_h if com_h else 0))
        return ns, cs, name_h, com_h, hh

    # 行が低くなるときは、見出しを詰めて行に高さを回す
    avail = a.h - room_t - room_b
    for compact in (False, True):
        ns, cs, name_h, com_h, hh = header(compact)
        row_h = min(1.0, (avail - hh - 0.1) / m)
        if row_h >= 0.5:
            break
    if row_h < 0.36:
        K.WARN.add("評価軸が多すぎて表が窮屈です。評価軸を減らすか、案の名前を短くしてください")
    # 表は本体の上端から（おすすめのラベルの上端が a.y）
    top = a.y + room_t
    body_top = top + hh + 0.1
    body_bottom = body_top + row_h * m
    sym_size = max(14, min(26, round(row_h * 72) / 2))  # 記号は行の高さに合わせる（飾りの大きさ）
    roomy = row_h >= 0.6
    txt_size = z["body"] if roomy else z["body_small"]
    lab_size = z["body_large"] if roomy else z["body"]

    # 1) おすすめ列の淡い面（強調の面は plain_card の色。線は最後に重ねる）
    fill, on = _strong(c)
    frame = None
    if rec is not None:
        fx = x0 + rec * (ow + gap)
        frame = Area(fx - ox, top - ov, ow + 2 * ox, body_bottom - top + 2 * ov)
        style.plain_card(s, frame, highlight=True).line.fill.background()
    # 2) 案の見出し（名前と一言の位置・大きさは全列でそろえる）
    ty = top + (hh - name_h - (0.06 + com_h if com_h else 0)) / 2
    for j, name in enumerate(names):
        x = x0 + j * (ow + gap)
        hl = j == rec
        box = style.plain_card(s, Area(x, top, ow, hh))
        if hl:
            _recolor(box, fill=fill, line=fill)
        style.text(s, x + 0.15, ty, tw, name_h, name, ns, on if hl else c["heading"], bold=True,
                   align="c", anchor="m", label="案の名前", line_spacing=1.0)
        com = comments[j] if j < len(comments) else ""
        if com:
            style.text(s, x + 0.15, ty + name_h + 0.06, tw, com_h + 0.02, com, cs,
                       on if hl else c["text_muted"], align="c", anchor="t", label="案の一言")
    # 3) 評価軸の行（最後の行の下にも線を引いて表を閉じる）
    for i, name in enumerate(crit):
        y = body_top + i * row_h
        style.text(s, a.x + 0.05, y, lw - 0.1, row_h, name, lab_size, c["heading"], bold=True,
                   anchor="m", fit_min=10, label="評価軸")
        style.row_rule(s, a.x, a.right, y + row_h)
        row = scores[i] if i < len(scores) else []
        for j in range(n):
            v = _s(row[j]).strip() if j < len(row) else ""
            if not v:
                continue
            x = x0 + j * (ow + gap)
            if v in SYMBOLS:
                style.text(s, x, y, ow, row_h, v, sym_size, c.get(SYMBOLS[v], c["text"]),
                           align="c", anchor="m")
            else:
                K.max_chars_warn(v, 16, "評価の文字")
                style.text(s, x + 0.12, y, ow - 0.24, row_h, v, txt_size, c["text"], align="c",
                           anchor="m", fit_min=9, label="評価の文字")
    # 4) おすすめの枠とラベル（ラベルは枠の上の線の上下中央）
    if frame:
        _outline(style, s, frame, fill, 2.0)
        label = "おすすめ"
        ps = z["body_small"]
        pw = _pill_w(label, ps, 1.2)
        px = frame.x + (frame.w - pw) / 2
        b = K.add_box(s.shapes, px, frame.y - pill_h / 2, pw, pill_h,
                      fill=fill, geom="pill")
        style.text(s, 0, 0, 0, 0, label, ps, on, bold=True, align="c",
                   anchor="m", shape=b, margin=0.03)


# ====================================================================== matrix（2×2）

def _axis(v):
    """軸をそろえる。文字だけなら軸の名前とみなす。"""
    if isinstance(v, str):
        return {"label": v, "low": "", "high": ""}
    if isinstance(v, dict):
        return {k: _s(v.get(k)) for k in ("label", "low", "high")}
    return {"label": "", "low": "", "high": ""}


def _quad(q):
    """象限をそろえる。文字だけなら名前、リストなら項目とみなす（位置は変えない）。"""
    if q is None or q == "":
        return {"title": "", "items": []}
    if isinstance(q, str):
        return {"title": q, "items": []}
    if isinstance(q, (list, tuple)):
        return {"title": "", "items": _text_list(q)}
    if not isinstance(q, dict):
        K.WARN.add("象限（quadrants）は {title, items} の形で書いてください")
        return {"title": "", "items": []}
    return {"title": _s(q.get("title")), "items": _text_list(q.get("items"))}


@register("matrix")
def matrix(style, sd):
    """2×2 のマトリクス。x / y: {label, low, high}、quadrants: 左上→右上→左下→右下の4つ
    （{title, items}）、highlight: 強調する象限（0〜3）。"""
    s, a = style.body(sd)
    c = style.c
    xa, ya = _axis(sd.get("x")), _axis(sd.get("y"))
    quads = [_quad(q) for q in _list(sd.get("quadrants"))]
    if len(quads) != 4:
        K.WARN.add(f"象限（quadrants）は 4 つにしてください（いまは {len(quads)} 個）")
    quads = (quads + [_quad(None)] * 4)[:4]
    hl = sd.get("highlight")
    if hl is not None and _index(hl, 4) is None:
        K.WARN.add("highlight は 0〜3 の数にしてください")
        hl = None
    for q in quads:
        K.max_chars_warn(q["title"], 14, "象限の名前")
        if len(q["items"]) > 4:
            K.WARN.add(f"象限の項目は 4 個までにしてください（「{q['title']}」は {len(q['items'])} 個）")
        for t in q["items"]:
            K.max_chars_warn(t, 24, "象限の項目")

    z = _scale(style)
    ts, ws = z["body"], z["caption"]      # 軸と象限の名前（ラベル）・軸の両端の言葉
    yw, xh = 1.0, 0.68                    # 縦軸の名前の列（意図した見出しの列）と、横軸の行
    gx, gy = a.x + yw, a.y
    gw, gh = a.right - gx, a.bottom - xh - gy
    gap = style.gutter()
    qw = (gw - gap) / 2
    axis = c["text_faint"]
    lx = a.x + 0.62                      # 縦軸の線
    ly = gy + gh + 0.3                   # 横軸の線
    # 縦軸の矢じりが本文の枠の上に出ないよう、少し下で止める
    _arrow_line(s, lx, ly, lx, gy + 0.06, axis, 1.75)
    _arrow_line(s, lx, ly, a.right, ly, axis, 1.75)
    fill, on = _strong(c)
    # 縦軸: 名前（縦書きのラベル）と、上端・下端の言葉
    if ya["label"]:
        K.max_chars_warn(ya["label"], 8, "縦軸の名前")
        vh = len(_vchars(ya["label"])) * ts * K.LINE_HEIGHT * VLS / 72
        ph = min(gh - 1.0, max(1.1, vh + 0.4))
        style.pill(s, lx - 0.2, gy + (gh - ph) / 2, 0.4, ph, "", "primary")
        _vtext(style, s, lx - 0.2, gy + (gh - ph) / 2, 0.4, ph, ya["label"], ts, on)
    for t, y in ((ya["high"], gy + 0.08), (ya["low"], gy + gh - 0.36)):
        if t:
            style.text(s, a.x, y, lx - a.x - 0.08, 0.36, t, ws, c["text_muted"], bold=True,
                       align="r", anchor="m", fit_min=9, label="軸の言葉")
    # 横軸: 名前（ラベル）と、左端・右端の言葉
    if xa["label"]:
        pw = _pill_w(xa["label"], ts, 1.1, gw / 2)
        style.pill(s, gx + (gw - pw) / 2, ly - _pill_h(ts) / 2, pw, _pill_h(ts), xa["label"],
                   "primary", ts)
    for t, x, al in ((xa["low"], gx, "l"), (xa["high"], a.right - 1.6, "r")):
        if t:
            style.text(s, x, ly + 0.05, 1.6, 0.32, t, ws, c["text_muted"], bold=True, align=al,
                       anchor="m", fit_min=9, label="軸の言葉")

    # 象限: 文字の大きさと列の数は4つでそろえる。名前（ラベル）と項目は同じ body
    pad = 0.24
    iw = qw - 2 * pad
    avail = gh - gap

    def qdims(compact):
        """象限の中の寸法（ラベルの文字・高さ・下の間、上下の余白、項目の間）。
        項目を小さくするときは、先にラベルと余白を詰める。"""
        if compact:
            return {"ps": z["body_small"], "ph": _pill_h(z["body_small"]), "pg": 0.08,
                    "top": 0.14, "bot": 0.1, "para": 3}
        return {"ps": ts, "ph": _pill_h(ts), "pg": 0.12, "top": 0.2, "bot": 0.16, "para": 4}

    def items_h(items, size, cols, para):
        if cols == 1:
            return _points_h(items, iw, size, para, 1.05)
        half = math.ceil(len(items) / 2)
        cw = (iw - 0.25) / 2
        return max(_points_h(items[:half], cw, size, para, 1.05),
                   _points_h(items[half:], cw, size, para, 1.05))

    def need(q, size, cols, compact):
        d = qdims(compact)
        return (d["top"] + d["bot"] + (d["ph"] + d["pg"] if q["title"] else 0)
                + items_h(q["items"], size, cols, d["para"]))

    def top_height(size, cols, compact):
        """収まるときは上段の高さ、収まらないときは None。
        上下が等分で収まれば等分。収まらなければ、文字を縮める前に中身の多い段を高くする。"""
        nt = max(need(quads[0], size, cols, compact), need(quads[1], size, cols, compact))
        nb = max(need(quads[2], size, cols, compact), need(quads[3], size, cols, compact))
        if max(nt, nb) <= avail / 2:
            return avail / 2
        if nt + nb <= avail:
            return nt + (avail - nt - nb) / 2
        return None

    # 1列で body → body_small → 2列（項目が3つ以上の象限があるとき）→ ラベルを縮めて 9pt まで
    multi = any(len(q["items"]) >= 3 for q in quads)
    col_opts = (1, 2) if multi else (1,)
    sizes = _uniq((z["body"], z["body_small"]))
    plans = [(v, k, False) for k in col_opts for v in sizes]
    plans += [(v, k, True) for v in _steps(z["body_small"] - 0.5, 9) for k in col_opts]
    for size, cols, compact in plans:
        ht = top_height(size, cols, compact)
        if ht is not None:
            break
    else:
        ht = avail / 2
        K.WARN.add("象限の項目が多すぎて枠に収まりません。項目を減らすか短くしてください")
    if ht is not None and size < 11:
        K.WARN.add(f"象限の項目が多いため文字が {size:g}pt まで小さくなっています。"
                   "項目を減らすか短くしてください")
    heights = [ht, avail - ht]
    d = qdims(compact)
    for i, q in enumerate(quads):
        r, col = divmod(i, 2)
        qa = Area(gx + col * (qw + gap), gy + r * (heights[0] + gap), qw, heights[r])
        on_hl = i == hl
        box = style.plain_card(s, qa, highlight=on_hl)
        if on_hl:
            _recolor(box, line=fill, line_w=2.0)
        y = qa.y + d["top"]
        if q["title"]:
            tone = "primary" if (hl is None or on_hl) else "soft"
            style.pill(s, qa.x + pad, y, _pill_w(q["title"], d["ps"], 1.0, iw), d["ph"], q["title"],
                       tone, d["ps"])
            y += d["ph"] + d["pg"]
        items = q["items"]
        if not items:
            continue
        half = math.ceil(len(items) / 2) if cols == 2 else len(items)
        cw = (iw - 0.25) / 2 if cols == 2 else iw
        for k, part in enumerate((items[:half], items[half:])):
            if part:
                style.text(s, qa.x + pad + k * (cw + 0.25), y, cw, qa.bottom - d["bot"] + 0.04 - y,
                           part, size, c["text"], anchor="t", fit_min=9, label="象限の項目",
                           bullet="・", bullet_color=fill, para_space=d["para"],
                           line_spacing=1.05)


# ====================================================================== process（手順・フェーズ）

def _step(st):
    """ステップをそろえる。文字（や数）だけならステップの名前とみなす。"""
    if isinstance(st, (str, int, float)) and not isinstance(st, bool):
        return {"title": _s(st), "period": "", "body": "", "points": []}
    if not isinstance(st, dict):
        return None
    title = st.get("title") if st.get("title") is not None else st.get("text")
    return {"title": _s(title), "period": _s(st.get("period")), "body": _s(st.get("body")),
            "points": _text_list(st.get("points"))}


@register("process")
def process(style, sd):
    """手順・フェーズ。steps: [{title, period, body, points}]（2〜6個）。
    上に矢羽根（期間。期間がなければ STEP 1, 2…）を並べ、その下に各ステップのカードを置く。"""
    s, a = style.body(sd)
    c = style.c
    raw = _list(sd.get("steps"))
    steps = [st for st in (_step(x) for x in raw) if st]
    if len(steps) < len(raw):
        K.WARN.add("ステップ（steps）は文字か {title, period, body, points} で書いてください")
    check_count(steps, 2, 6, "ステップ（steps）")
    if not steps:
        return
    n = len(steps)
    has_period = any(st["period"] for st in steps)
    body_max = 85 if n <= 3 else (60 if n == 4 else (45 if n == 5 else 30))
    points_max = 3 if n <= 4 else 2
    for st in steps:
        K.max_chars_warn(st["title"], 16, "ステップの名前")
        K.max_chars_warn(st["period"], 12, "ステップの期間")
        K.max_chars_warn(st["body"], body_max, "ステップの本文")
        if len(st["points"]) > points_max:
            K.WARN.add(f"ステップの箇条書きは {points_max} 個までにしてください（「{st['title']}」は "
                       f"{len(st['points'])} 個）")
    z = _scale(style)
    gap = style.gutter()
    w = (a.w - gap * (n - 1)) / n
    fill, on = _strong(c)

    def color(i):
        return fill
    # 文字の大きさ: 名前は item_head、本文と箇条書きは同じ大きさ（ステップが3つまでなら body_large）
    tsize = z["item_head"]
    first = z["body_large"] if n <= 3 else z["body"]
    wsize = z["body"] if n <= 4 else z["body_small"]   # 矢羽根の言葉

    def dims(compact):
        """矢羽根の高さ・カードの余白・ラベルの高さ・区切りの間。compact は詰めた組み方。"""
        if compact:
            return {"ch": 0.46, "cg": 0.14, "pad": 0.12, "lh": 0.22 if has_period else 0,
                    "g1": 0.06, "g2": 0.08}
        return {"ch": 0.56, "cg": 0.2, "pad": 0.2 if n <= 4 else 0.16,
                "lh": 0.28 if has_period else 0, "g1": 0.1, "g2": 0.12}

    def title_h(d):
        return max(_lines_h(st["title"], w - 2 * d["pad"], tsize, True, 2) for st in steps) + 0.04

    def body_h(st, size, d):
        if not st["body"]:
            return 0
        return _lines_h(st["body"], w - 2 * d["pad"], size, line_spacing=1.08) + 0.06

    def points_h(st, size, d):
        return _points_h(st["points"], w - 2 * d["pad"], size, 2, 1.0)

    def need(st, size, d):
        h = 2 * d["pad"] + d["lh"] + title_h(d) + d["g1"] + d["g2"]
        if st["body"]:
            h += body_h(st, size, d) + 0.04
        return h + points_h(st, size, d)

    # いちばん中身の多いカードで大きさを決め、全部のカードを同じ大きさで書く。
    # body_small でも収まらなければ、余白を詰めて 9pt まで
    plans = [(v, False) for v in _uniq((first, z["body"], z["body_small"])) if v <= first]
    plans += [(v, True) for v in _steps(z["body_small"] - 0.5, 9)]
    over = False
    for size, compact in plans:
        d = dims(compact)
        if max(need(st, size, d) for st in steps) <= a.h - d["ch"] - d["cg"]:
            break
    else:
        over = True
        K.WARN.add("ステップの文字が多すぎて枠に収まりません。本文を短くするか、箇条書きを減らしてください")
    ch, pad = d["ch"], d["pad"]
    iw = w - 2 * pad
    th = title_h(d)
    avail = a.h - ch - d["cg"]
    # カードは中身に合わせた高さ（低すぎると貧弱なので 1.6 インチ以上）。
    # 矢羽根は本体の上端から置き、余りは下に残す
    h = min(avail, max(max(need(st, size, d) for st in steps) + 0.1, min(avail, 1.6)))
    y0 = a.y
    # 矢羽根の列（1つ目は五角形、2つ目からは山形）。どの矢羽根もその下のカードと同じ x・幅。
    # 期間があるときは期間だけを書き（STEP n はカードのラベル）、ないときは STEP n を書く
    tip = ch * 0.42
    for i, st in enumerate(steps):
        x = a.x + i * (w + gap)
        shp = K.add_box(s.shapes, x, y0, w, ch, fill=color(i),
                        geom="pentagon" if i == 0 else "chevron")
        shp.adjustments[0] = tip / ch
        word = st["period"] if has_period else f"STEP {i + 1}"
        if word:
            # 文字の枠は左右に同じだけ空けて、カードの中央にそろえる
            style.text(s, x + tip, y0, w - 2 * tip, ch, word, wsize, on, bold=True, align="c",
                       anchor="m", fit_min=9, label="ステップの期間", margin=0)
    # カード
    top = y0 + ch + d["cg"]
    fit = 9 if over else None             # 収まるときは縮めない（大きさをそろえる）
    for i, st in enumerate(steps):
        x = a.x + i * (w + gap)
        card = style.plain_card(s, Area(x, top, w, h))
        bg = _fill_of(card, c["card_bg"])
        y = top + pad
        bottom = top + h - pad + 0.06
        if has_period:
            style.text(s, x + pad, y - 0.02, iw, d["lh"], f"STEP {i + 1}", z["caption"],
                       _ink(c, color(i), bg), bold=True, anchor="m",
                       latin=style.font("number_latin"), spacing=1)
            y += d["lh"]
        style.text(s, x + pad, y, iw, th, st["title"], tsize, c["heading"], bold=True,
                   anchor="m", fit_min=11, label="ステップの名前", line_spacing=1.0)
        y += th + d["g1"]
        _rule(style, s, x + pad, x + w - pad, y, bg)
        y += d["g2"]
        if st["body"]:
            bh = body_h(st, size, d)
            style.text(s, x + pad, y, iw, min(bh, bottom - y), st["body"], size, c["text"],
                       anchor="t", fit_min=fit, label="ステップの本文", line_spacing=1.08)
            y += bh + 0.04
        if st["points"]:
            ph = points_h(st, size, d) + 0.06
            style.text(s, x + pad, y, iw, max(0.2, min(ph, bottom - y)), st["points"], size,
                       c["text"], anchor="t", fit_min=fit, label="ステップの箇条書き", bullet="・",
                       bullet_color=color(i), para_space=2)


# ====================================================================== tree（ロジックツリー）

def _branch(b):
    """枝をそろえる。文字だけなら葉のない枝。正しくない書き方は None。"""
    if isinstance(b, str):
        return {"text": b, "children": []}
    if isinstance(b, dict):
        text = b.get("text") if b.get("text") is not None else b.get("title")
        return {"text": _s(text), "children": _text_list(b.get("children"))}
    return None


@register("tree")
def tree(style, sd):
    """ロジックツリー（左から右へ分解）。root: 根、branches: [{text, children}]（2〜4個、
    children は各 0〜4 個）、highlight: 強調する枝（任意）。
    葉が多くて1つずつの箱に入らないときは、枝ごとに葉を1枚のカードの箇条書きにまとめる。"""
    s, a = style.body(sd)
    c = style.c
    root = _s(sd.get("root"))
    raw = _list(sd.get("branches"))
    branches = [b for b in (_branch(x) for x in raw) if b]
    if len(branches) < len(raw):
        K.WARN.add("枝（branches）は文字か {text, children} で書いてください")
    check_count(branches, 2, 4, "枝（branches）")
    if not branches:
        return
    if not root:
        K.WARN.add("tree には root（根）が必要です")
    K.max_chars_warn(root, 20, "ツリーの根")
    for b in branches:
        K.max_chars_warn(b["text"], 16, "ツリーの枝")
        if len(b["children"]) > 4:
            K.WARN.add(f"枝の子は 4 個までにしてください（「{b['text']}」は {len(b['children'])} 個）")
        for t in b["children"]:
            K.max_chars_warn(t, 22, "ツリーの葉")
    hl = sd.get("highlight")
    if hl is not None and _index(hl, len(branches)) is None:
        K.WARN.add(f"highlight は 0〜{len(branches) - 1} の数にしてください")
        hl = None
    nb = len(branches)
    has_child = any(b["children"] for b in branches)
    kid_n = [len(b["children"]) for b in branches]
    z = _scale(style)
    hs = z["item_head"]                   # 根と枝の文字
    leaf_sizes = _uniq((z["body"], z["body_small"]))
    line = c.get("text_faint", c["text_muted"])   # 線は白地と 3:1 ほどの差のある色（淡すぎると見えない）
    fill, on = _strong(c)
    lw_n, lw_h = 1.5, 2.0                 # 線の太さ（ふつう・強調）
    # 横の配置: 根は a.x から、葉（葉がなければ枝）は a.right まで。
    # 根と枝の箱は文字が1行に入る幅（上限あり）にし、残りを葉に回す
    rw = min(2.9, max(2.0, _em(root, True) * hs / 72 / SAFE + 0.5))
    b_em = max(_em(b["text"], True) for b in branches)
    bw = min(3.2, max(2.4, b_em * hs / 72 / SAFE + 0.5))
    gx = 0.7                              # 列の間（線を引く場所）
    leaves = [t for b in branches for t in b["children"]]
    leaf_em = max([_em(t) for t in leaves] or [0])

    def leaf_room(rw, bw, gx):
        return a.right - (a.x + rw + gx + bw + gx)

    def branch_size(bw):
        """枝の文字の大きさ（全部の枝でそろえる）。1行に収まるなら body_large まで縮めて1行に。"""
        return next((v for v in _uniq((hs, z["body_large"]))
                     if all(K.count_lines(K.plain(b["text"]), (bw - 0.36) * SAFE, v, True) <= 1
                            for b in branches)), hs)

    # 1) 葉を1つずつ箱に入れる組み方。葉が body_small でも箱に入らないときは使わない
    box = None
    if has_child:
        lw = leaf_room(rw, bw, gx)
        for gg, sg, lo in ((0.24, 0.1, 0.46), (0.16, 0.06, 0.39)):
            slots = [max(1, k) for k in kid_n]
            slot_h = min(0.78, (a.h - gg * (nb - 1) - sg * (sum(slots) - nb)) / sum(slots))
            if slot_h < lo:
                continue
            fs = next((v for v in leaf_sizes
                       if max(K.count_lines(K.plain(t), (lw - 0.32) * SAFE, v) for t in leaves)
                       * v * K.LINE_HEIGHT / 72 + 0.12 <= slot_h), None)
            if fs is not None:
                box = (gg, sg, slot_h, fs)
                break
    # 2) 枝ごとに葉を1枚のカード（箇条書き）にまとめる組み方。1列 → 2列 → 縮める
    card = None
    if has_child and box is None:
        gg = 0.16

        def card_need(kids, size, cols, cw):
            if not kids:
                return 0.0
            if cols == 1:
                return _points_h(kids, cw - 0.36, size, 3, 1.0) + 0.24
            half = math.ceil(len(kids) / 2)
            col_w = (cw - 0.36 - 0.25) / 2
            return max(_points_h(kids[:half], col_w, size, 3, 1.0),
                       _points_h(kids[half:], col_w, size, 3, 1.0)) + 0.24

        def try_cards(rw, bw, gx):
            cw = leaf_room(rw, bw, gx)
            bh = max(0.56, _lines_h(max((b["text"] for b in branches), key=len), bw - 0.36,
                                    branch_size(bw), True) + 0.24)
            plans = [(v, k) for k in (1, 2) for v in leaf_sizes]
            plans += [(v, k) for v in _steps(z["body_small"] - 0.5, 10) for k in (1, 2)]
            for size, cols in plans:
                needs = [max(bh, card_need(b["children"], size, cols, cw)) for b in branches]
                if sum(needs) + gg * (nb - 1) <= a.h:
                    return size, cols, cw, needs, bh
            return None

        card = try_cards(rw, bw, gx)
        if card is None:
            # 根と枝の箱を細くして（枝は body_large で1行に入る幅までは残す）葉のカードの幅をとる
            rw, gx = min(rw, 2.0), 0.45
            bw = min(bw, max(2.3, min(2.8, b_em * z["body_large"] / 72 / SAFE + 0.4)))
            card = try_cards(rw, bw, gx)
        if card is None:
            K.WARN.add("ツリーの葉が多すぎて枠に収まりません。葉を減らすか、"
                       "枝ごとに3つ以上あるときは葉を 16 文字以内にしてください")
            size, cols, cw = 10, 2, leaf_room(rw, bw, gx)
            needs = [max(0.56, card_need(b["children"], size, cols, cw)) for b in branches]
            card = (size, cols, cw, needs, 0.56)
    bs = branch_size(bw)
    # 縦の大きさ: 枝の箱は全部同じ高さ。葉の束（または葉のカード）の上下中央に置く
    if card:
        size, cols, cw, needs, bh = card
        spare = a.h - sum(needs) - gg * (nb - 1)
        if spare >= 0:
            group_h = [h + min(0.2, spare / nb) for h in needs]
        else:
            # 収まらないときも枠からは出さない（文字は fit_min で縮む）
            group_h = [h * (a.h - gg * (nb - 1)) / sum(needs) for h in needs]
        bh = min([bh] + group_h)
    elif box:
        gg, sg, slot_h, fs = box
        kids_h = [max(1, k) * slot_h + (max(1, k) - 1) * sg for k in kid_n]
        bh = max(slot_h, 0.9)
        while bh > slot_h and sum(max(h, bh) for h in kids_h) + gg * (nb - 1) > a.h:
            bh = max(slot_h, bh - 0.02)
        group_h = [max(h, bh) for h in kids_h]
    else:
        gg = 0.24
        bh = min(1.0, (a.h - gg * (nb - 1)) / nb)
        group_h = [bh] * nb
    rx = a.x
    bx = rx + rw + gx
    if not has_child:
        bw = a.right - bx
    lx = bx + bw + gx
    leaf_w = a.right - lx
    sx = bx + bw + gx / 2                 # 枝から葉への線を分ける位置
    jx = rx + rw + gx / 2                 # 根から枝への線を分ける位置（幹）
    # 根の高さと縦の位置: 根の中心を強調する枝（なければ真ん中の枝）の中心にそろえて、
    # 根から枝への線を1本のまっすぐな線にする。本体は a.y から（根が上に出るときだけ枝を下げる）
    used = sum(group_h) + gg * (nb - 1)
    rel = [sum(group_h[:i]) + gg * i + group_h[i] / 2 for i in range(nb)]
    if hl is not None:
        target = rel[hl]
    elif nb % 2:
        target = rel[nb // 2]
    else:
        target = (rel[0] + rel[-1]) / 2
    # 根の文字は 1 行に収まるなら body_large まで縮めて 1 行に（1〜2 文字だけ次の行に送らない）
    rs = next((v for v in _uniq((hs, z["body_large"]))
               if K.count_lines(K.plain(root), (rw - 0.4) * SAFE, v, True) <= 1), hs)
    rh = max(1.1, _lines_h(root, rw - 0.4, rs, True) + 0.5)
    rh = max(0.5, min(rh, 2 * (a.h - used + target), 2 * (a.h - target), a.h))
    top = a.y + max(0.0, rh / 2 - target)
    rm = top + target
    y = top
    mids = []
    for i, b in enumerate(branches):
        gh = group_h[i]
        on_hl = i == hl
        lc, lwid = (fill, lw_h) if on_hl else (line, lw_n)
        mid = y + gh / 2
        mids.append(mid)
        bbox = style.plain_card(s, Area(bx, mid - bh / 2, bw, bh), idx=i, highlight=True)
        if on_hl:
            _recolor(bbox, fill=fill, line=fill)
        style.text(s, bx + 0.18, mid - bh / 2, _balanced_w(b["text"], bw - 0.36, bs, True), bh,
                   b["text"], bs, on if on_hl else c["heading"], bold=True, anchor="m",
                   fit_min=10, label="ツリーの枝")
        kids = b["children"]
        if kids and card:
            # 葉をまとめたカード（2列のときは上から下へ、左の列から埋める）
            kb = style.plain_card(s, Area(lx, y, leaf_w, gh))
            if on_hl:
                _recolor(kb, line=fill, line_w=1.5)
            half = math.ceil(len(kids) / 2) if cols == 2 else len(kids)
            col_w = (leaf_w - 0.36 - 0.25) / 2 if cols == 2 else leaf_w - 0.36
            for k, part in enumerate((kids[:half], kids[half:])):
                if part:
                    style.text(s, lx + 0.18 + k * (col_w + 0.25), y + 0.12, col_w, gh - 0.24, part,
                               size, c["text"], anchor="m", fit_min=9, label="ツリーの葉",
                               bullet="・", bullet_color=lc, para_space=3)
            K.add_line(s.shapes, bx + bw, mid, lx, mid, lc, lwid)
        elif kids:
            # 葉の束は枝の中心にそろえる（葉が1つなら枝からまっすぐ1本の線）
            ky0 = mid - (len(kids) * slot_h + (len(kids) - 1) * sg) / 2
            kid_mids = []
            for k, t in enumerate(kids):
                ky = ky0 + k * (slot_h + sg)
                kb = style.plain_card(s, Area(lx, ky, leaf_w, slot_h))
                if on_hl:
                    _recolor(kb, line=fill, line_w=1.5)
                style.text(s, lx + 0.16, ky, leaf_w - 0.32, slot_h, t, fs, c["text"],
                           anchor="m", fit_min=9, label="ツリーの葉")
                kid_mids.append(ky + slot_h / 2)
            if len(kids) == 1:
                K.add_line(s.shapes, bx + bw, mid, lx, mid, lc, lwid)
            else:
                K.add_line(s.shapes, bx + bw, mid, sx, mid, lc, lwid)
                K.add_line(s.shapes, sx, kid_mids[0], sx, kid_mids[-1], lc, lwid)
                for km in kid_mids:
                    K.add_line(s.shapes, sx, km, lx, km, lc, lwid)
        y += gh + gg
    # 根と、根から枝への線（幹を1本引き、各枝へ分ける。強調する枝へは根から1本の濃い線）
    rbox = style.plain_card(s, Area(rx, rm - rh / 2, rw, rh))
    _recolor(rbox, fill=fill, line=fill)
    tw = _balanced_w(root, rw - 0.4, rs, True)
    style.text(s, rx + (rw - tw) / 2, rm - rh / 2, tw, rh, root, rs, on, bold=True, align="c",
               anchor="m", fit_min=11, label="ツリーの根")
    K.add_line(s.shapes, jx, mids[0], jx, mids[-1], line, lw_n)
    straight = None                       # 根からまっすぐつながる枝
    for i, m in enumerate(mids):
        if abs(m - rm) < 1e-6 and straight is None:
            straight = i
    for i, m in enumerate(mids):
        if i == hl:
            continue
        K.add_line(s.shapes, rx + rw if i == straight else jx, m, bx, m, line, lw_n)
    if hl is not None:
        K.add_line(s.shapes, rx + rw, rm, bx, rm, fill, lw_h)
    elif straight is None:
        K.add_line(s.shapes, rx + rw, rm, jx, rm, line, lw_n)
