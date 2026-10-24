"""スライドの型ごとの組み立て。

各関数は (style, sd) を受け取る。sd は deck.json の slides の1要素。
見た目の違いは style（styles_impl.py）に任せ、ここでは配置だけを決める。
"""
import re

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import slidekit as K
from styles_impl import Area

REGISTRY = {}


def register(name):
    def deco(fn):
        REGISTRY[name] = fn
        return fn
    return deco


def as_items(items):
    """文字列だけの項目も {text: ...} にそろえる。
    リストの代わりに文字列1つ・オブジェクト1つが来ても、1項目として扱う。"""
    if items is None:
        return []
    if isinstance(items, (str, dict, int, float)):
        items = [items]
    out = []
    for it in items:
        if it is None:
            continue
        if isinstance(it, dict):
            out.append(dict(it))
        else:
            out.append({"text": K.as_text(it)})
    return out


def check_count(items, lo, hi, what):
    n = len(items)
    if n < lo or n > hi:
        K.WARN.add(f"{what}は {lo}〜{hi} 個にしてください（いまは {n} 個）")


# ====================================================================== 構造

@register("cover")
def cover(style, sd):
    style.cover(sd)


@register("section")
def section(style, sd):
    style.section(sd)


@register("closing")
def closing(style, sd):
    style.closing(sd)


# ====================================================================== 箇条書き

@register("bullets")
def bullets(style, sd):
    """要点＋根拠。items: [{text, detail}] または文字列のリスト（1〜6個）。"""
    s, a = style.body(sd)
    c = style.c
    items = as_items(sd.get("items"))
    check_count(items, 1, 6, "箇条書きの項目")
    if not items:
        return
    n = len(items)
    has_detail = any(it.get("detail") for it in items)
    d = style.MARKER["m"][0]
    tx = a.x + d + 0.25
    tw = a.w - d - 0.25
    head_size = style.size("list_head", 17) if has_detail else style.size("list_head", 17) + 1
    det_size = style.size("list_detail", 13)
    # 各項目の高さを見積もり、項目の間隔（ピッチ）をそろえる
    blocks = []
    for it in items:
        K.max_chars_warn(it.get("text", ""), 40, "箇条書きの要点")
        hh = max(0.42, K.text_height(K.plain(it.get("text", "")), tw, head_size, True) + 0.06)
        dh = 0.0
        if it.get("detail"):
            K.max_chars_warn(it["detail"], 90 if n <= 3 else (75 if n <= 5 else 35), "要点の補足")
            dh = K.text_height(K.plain(it["detail"]), tw, det_size, line_spacing=1.05) + 0.06
        blocks.append((hh, dh))
    need = max(h + dd for h, dd in blocks)
    pitch = min(a.h / n, need + (0.62 if has_detail else 0.5))
    pitch = max(pitch, min(a.h / n, need + 0.2))
    y = a.y
    for i, (it, (hh, dh)) in enumerate(zip(items, blocks)):
        style.marker(s, a.x, y + (hh - d) / 2, it.get("mark", i + 1), "m")
        if has_detail:
            style.text(s, tx, y, tw, hh, it.get("text", ""), head_size, c["heading"], bold=True,
                       anchor="m", fit_min=12, label="要点")
            if it.get("detail"):
                style.text(s, tx, y + hh, tw, max(dh, pitch - hh - 0.2), it["detail"], det_size,
                           c["text"], anchor="t", fit_min=10, label="要点の補足", line_spacing=1.05)
        else:
            style.text(s, tx, y, tw, hh, it.get("text", ""), head_size, c["text"], anchor="m",
                       fit_min=12, label="要点")
        if i < n - 1:
            # 区切り線は、この項目の下端と次の項目の上端のちょうど中間に引く
            gap_mid = y + (hh + dh + pitch) / 2
            style.row_rule(s, a.x, a.right, gap_mid)
        y += pitch


# ====================================================================== 数字の強調

@register("kpi")
def kpi(style, sd):
    """数字の強調。items: [{label, value, unit, note}]（1〜4個）。points で下に補足を置ける。"""
    s, a = style.body(sd)
    c = style.c
    items = as_items(sd.get("items"))
    check_count(items, 1, 4, "数字の項目")
    if not items:
        return
    n = len(items)
    gap = style.gutter()
    points = sd.get("points") or []
    tile_h = min(a.h, 3.2) if not points else min(a.h * 0.62, 2.8)
    tile_w = (a.w - gap * (n - 1)) / n
    if n == 1:
        tile_w = min(a.w, 6.0)
    x0 = a.x + (a.w - (tile_w * n + gap * (n - 1))) / 2
    y0 = a.y
    for i, it in enumerate(items):
        K.max_chars_warn(it.get("label", ""), 12, "数字の項目名")
        x = x0 + i * (tile_w + gap)
        inner = style.kpi_tile(s, Area(x, y0, tile_w, tile_h), it, i)
        note = it.get("note", "")
        note_h = 0.5 if note else 0
        num_h = inner.h - note_h - 0.2
        style.kpi_value(s, Area(inner.x + 0.15, inner.y + 0.05, inner.w - 0.3, num_h),
                        it.get("value", ""), it.get("unit", ""),
                        style.size("kpi_number", [80, 72, 62, 52])[min(n, 4) - 1])
        if note:
            K.max_chars_warn(note, 40, "数字の補足")
            style.kpi_note(s, Area(inner.x, inner.bottom - note_h - 0.12, inner.w, note_h), note)
    if points:
        py = y0 + tile_h + 0.3
        ph = a.bottom - py
        style.text(s, a.x + 0.1, py, a.w - 0.2, ph, [K.plain(p) if False else p for p in points],
                   14, c["text"], anchor="t", fit_min=10, label="補足の箇条書き", bullet="●",
                   bullet_color=c["primary"], para_space=4)


# ====================================================================== カード

def grid_rows(n):
    """n 枚のカードの段組み。5枚は上2・下3。"""
    if n <= 3:
        return [n]
    if n == 4:
        return [2, 2]
    if n == 5:
        return [2, 3]
    return [3, n - 3]


def card_grid(n, a, gap=0.28, needed=None, min_h=1.9):
    """n 枚のカードの配置を返す。needed（各カードに必要な高さ）を渡すと、
    中身が少ないときはカードを縮め、全体を上下中央に寄せる。"""
    rows = grid_rows(n)
    nrows = len(rows)
    full_h = (a.h - gap * (nrows - 1)) / nrows
    heights = []
    k = 0
    for cols in rows:
        if needed:
            need = max(needed[k:k + cols])
            heights.append(min(full_h, max(min_h, need)))
        else:
            heights.append(full_h)
        k += cols
    # カードは例外として、縮めたときは全体を上下中央に寄せる（下に大きな空きを残さない）。
    # 中身で領域が埋まるときは上詰めと同じ位置になる
    total = sum(heights) + gap * (nrows - 1)
    y = a.y + max(0.0, (a.h - total) / 2)
    boxes = []
    for cols, rh in zip(rows, heights):
        w = (a.w - gap * (cols - 1)) / cols
        for j in range(cols):
            boxes.append(Area(a.x + j * (w + gap), y, w, rh))
        y += rh + gap
    return boxes


def card_needed_height(style, it, w, size=13):
    """カード1枚に必要な高さの見積もり。"""
    inner_w = w - 0.44
    h = style.card_header_h(it, w)
    body = K.plain(it.get("body", ""))
    points = it.get("points") or []
    if body:
        h += K.text_height(body, inner_w, size, line_spacing=1.08) + 0.1
    if points:
        h += K.text_height("\n".join(K.plain(p) for p in points), inner_w - 0.22, size - 0.5) + 0.1
    return h + 0.3


@register("cards")
def cards(style, sd):
    """並列の項目をカードで並べる。items: [{label, heading, body, points}]（2〜6個）。"""
    s, a = style.body(sd)
    c = style.c
    items = as_items(sd.get("items"))
    check_count(items, 2, 6, "カード")
    if not items:
        return
    rows = grid_rows(len(items))
    widths = []
    for cols in rows:
        widths += [(a.w - 0.28 * (cols - 1)) / cols] * cols
    # 余裕があれば本文を大きく（16pt）、なければ 13pt
    full_h = (a.h - 0.28 * (len(rows) - 1)) / len(rows)
    size = style.size("card_body_large", 16) if len(items) <= 3 else style.size("card_body", 14)
    needed = [card_needed_height(style, it, w, size) for it, w in zip(items, widths)]
    if max(needed) > full_h:
        size = 13
        needed = [card_needed_height(style, it, w, size) for it, w in zip(items, widths)]
    # カードは中身に合わせた高さ（少し余白を足す）。低すぎると貧弱なので 1.6 インチ以上
    boxes = card_grid(len(items), a, needed=[n + 0.25 for n in needed],
                      min_h=min(full_h, 1.6))
    # 5枚（上2・下3）のときは、上の2枚を目立たせる（highlight で指定もできる）
    highlight = sd.get("highlight")
    if highlight is None:
        highlight = [0, 1] if len(items) == 5 else []
    if isinstance(highlight, int):
        highlight = [highlight]
    for i, (it, box) in enumerate(zip(items, boxes)):
        K.max_chars_warn(it.get("heading", ""), 24, "カードの見出し")
        inner = style.card(s, box, it, i, highlight=i in highlight)
        body = it.get("body", "")
        points = it.get("points") or []
        if body:
            K.max_chars_warn(body, {1: 110, 2: 110, 3: 100, 4: 55}.get(len(items), 30), "カードの本文")
        if body and points:
            body_h = min(inner.h * 0.6, K.text_height(K.plain(body), inner.w, size, line_spacing=1.08) + 0.12)
            style.text(s, inner.x, inner.y, inner.w, body_h, body, size, c["text"],
                       anchor="t", fit_min=9.5, label="カードの本文", line_spacing=1.08)
            py = inner.y + body_h + 0.06
            style.text(s, inner.x, py, inner.w, inner.bottom - py, points, size - 0.5, c["text"],
                       anchor="t", fit_min=9, label="カードの箇条書き", bullet="・",
                       bullet_color=c["primary"], para_space=2)
        elif body:
            style.text(s, inner.x, inner.y, inner.w, inner.h, body, size, c["text"],
                       anchor="t", fit_min=9.5, label="カードの本文", line_spacing=1.08)
        elif points:
            style.text(s, inner.x, inner.y, inner.w, inner.h, points, size - 0.5, c["text"],
                       anchor="t", fit_min=9, label="カードの箇条書き", bullet="・",
                       bullet_color=c["primary"], para_space=2)


# ====================================================================== 表

def _cell_border(cell, color, width_pt=0.75, sides=("L", "R", "T", "B")):
    tcPr = cell._tc.get_or_add_tcPr()
    for side in ("L", "R", "T", "B"):
        tag = qn(f"a:ln{side}")
        old = tcPr.find(tag)
        if old is not None:
            tcPr.remove(old)
    # 線の指定は塗りより前に置く
    insert_at = 0
    for side in ("L", "R", "T", "B"):
        ln = etree.Element(qn(f"a:ln{side}"), w=str(int(width_pt * 12700)), cap="flat",
                           cmpd="sng", algn="ctr")
        if side in sides:
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=color)
        else:
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(insert_at, ln)
        insert_at += 1


def _col_widths(header, rows, total, given=None):
    ncol = len(header)
    if given and len(given) == ncol:
        s = float(sum(given))
        return [total * g / s for g in given]
    weights = []
    for j in range(ncol):
        texts = [K.plain(str(header[j]))] + [K.plain(str(r[j])) if j < len(r) else "" for r in rows]
        longest = max(K.text_em(t) for t in texts)
        weights.append(min(max(longest, 3.0), 22.0))
    s = sum(weights)
    return [total * w / s for w in weights]


DATE = re.compile(r"^(\d{2,4}[/.\-年]\d{1,2}([/.\-月]\d{1,2}日?)?|\d{1,2}/\d{1,2}|\d{1,2}月(\d{1,2}日)?)$")
NUMERIC = re.compile(r"^[+\-−±]?[\d,，.]+\s*[%％件人回分秒点円倍個台日年月週名社]*"
                     r"(時間|万円|億円|千円|か月|ヶ月|ポイント|pt)?$")


def render_table(style, s, a, header, rows, col_widths=None, aligns=None, emph_col=None,
                 cell_style=None, size_max=13.0, size_min=9.0, grow=True, label="表",
                 highlight_rows=()):
    """表を描く（ほかの型からも使う）。

    header: 列名のリスト、rows: 行のリスト。
    aligns: {列番号: "l"/"c"/"r"}。emph_col: 太字にする列。
    cell_style(i, j, text) が dict(fill=, color=, bold=, align=) を返すと、そのセルだけ上書きする
    （i は本文の行番号 0〜、j は列番号）。
    戻り値: 実際に使った高さ（インチ）。
    """
    tc = style.table_colors()
    header = [str(h) for h in header]
    rows = [[str(v) for v in r] for r in rows]
    ncol = len(header)
    widths = _col_widths(header, rows, a.w, col_widths)
    aligns = {int(k): v for k, v in (aligns or {}).items()}
    # 列の寄せ（指定がなければ自動）: 数字だけの列は右、日付や短い記号の列は中央
    for j in range(ncol):
        if j in aligns or not rows:
            continue
        cells = [K.plain(r[j]) if j < len(r) else "" for r in rows]
        if j > 0 and all(NUMERIC.match(t) for t in cells):
            aligns[j] = "r"
        elif all(DATE.match(t) or len(t) <= 2 for t in cells):
            aligns[j] = "c"
    pad = 0.12

    def needed(size):
        total = 0
        heights = []
        for r_i, r in enumerate([header] + rows):
            lines = 1
            for j in range(ncol):
                t = K.plain(r[j]) if j < len(r) else ""
                lines = max(lines, K.count_lines(t, widths[j] - 2 * pad, size, r_i == 0))
            h = lines * size * K.LINE_HEIGHT / 72 + 0.14
            heights.append(max(h, 0.38))
            total += heights[-1]
        return total, heights

    size = min(size_max, style.size("table_body", size_max))
    while size > size_min:
        total, heights = needed(size)
        if total <= a.h:
            break
        size -= 0.5
    total, heights = needed(size)
    if total > a.h + 0.05:
        K.WARN.add(f"{label}が枠に収まりません。行を減らすか、文を短くしてください")
    elif grow and len(heights) > 1:
        # 余白が大きいときは本文の行を広げて読みやすくする（1行 0.62 インチまで。見出し行はそのまま）
        spare = a.h - total
        g = min(spare / (len(heights) - 1), max(0.0, 0.62 - min(heights[1:])))
        if g > 0:
            heights = [heights[0]] + [h + g for h in heights[1:]]
            total = sum(heights)
    gf = s.shapes.add_table(len(rows) + 1, ncol, Inches(a.x), Inches(a.y), Inches(a.w),
                            Inches(min(total, a.h)))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # 「スタイルなし、表のグリッドなし」（色と線は下で指定）
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, h in enumerate(heights):
        tbl.rows[i].height = Inches(h)
    for i, r in enumerate([header] + rows):
        for j in range(ncol):
            cell = tbl.cell(i, j)
            text = r[j] if j < len(r) else ""
            cell.margin_left = cell.margin_right = Inches(pad)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = K.ANCHOR["m"]
            if i == 0:
                bg, color, bold = tc["head_bg"], tc["head_text"], True
                # 見出しの寄せは、その列の中身の寄せにそろえる
                align = aligns.get(j, "l")
                fsize = min(size, style.size("table_head", size))
            else:
                bg = tc["row_alt"] if i % 2 == 0 else tc["row_bg"]
                color = tc["first_col"] if j == 0 else tc["text"]
                bold = j == 0 or (emph_col is not None and j == emph_col)
                if emph_col is not None and j == emph_col:
                    color = tc["first_col"]
                align = aligns.get(j, "l")
                fsize = size
                if (i - 1) in highlight_rows:
                    bg = style.c.get("tint", style.c["bg_alt"]) if j else bg
                    bold = True
                    color = tc["first_col"]
                if cell_style:
                    over = cell_style(i - 1, j, text) or {}
                    bg = over.get("fill", bg)
                    color = over.get("color", color)
                    bold = over.get("bold", bold)
                    align = over.get("align", align)
            cell.fill.solid()
            cell.fill.fore_color.rgb = K.rgb(bg)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            K.set_paragraph(p, align, 1.0)
            K.write_runs(p, text, fsize, color, bold=bold, emph_color=style.c.get("emph"),
                         font=style.font("ja"), latin=style.font("latin"))
            _cell_border(cell, tc["line"], 0.75, ("L", "R", "T", "B"))
    return min(total, a.h)


@register("table")
def table(style, sd):
    """表。header: [列名], rows: [[...]]。col_widths で列幅の比率、emphasis_col で強調列、
    align で列ごとの寄せ（{"1": "r"} など）。"""
    # 注記は表のすぐ下に置くので、下に確保させない
    note = sd.get("note")
    sd = {k: v for k, v in sd.items() if k != "note"}
    s, a = style.body(sd)
    if note:
        a = Area(a.x, a.y, a.w, a.h - 0.5)
    header = sd.get("header", [])
    rows = sd.get("rows", [])
    if not header or not rows:
        K.WARN.add("表には header と rows が必要です")
        return
    if len(rows) > 10:
        K.WARN.add(f"表の行が {len(rows)} 行あります（10 行以内を推奨。分けるか要約してください）")
    if len(header) > 7:
        K.WARN.add(f"表の列が {len(header)} 列あります（7 列以内を推奨）")
    hl = sd.get("highlight_row")
    hl = [hl] if isinstance(hl, int) else (hl or [])
    used = render_table(style, s, a, header, rows, sd.get("col_widths"), sd.get("align"),
                        sd.get("emphasis_col"), highlight_rows=hl)
    if note:
        K.max_chars_warn(note, 60, "注記")
        style.note(s, Area(a.x, a.y + used + 0.12, a.w, 0.34), note)


# ====================================================================== 追加の型
# types_*.py に書いた型をここで読み込む（ファイルを足すだけで型が増える）。
def _load_type_modules():
    import importlib
    from pathlib import Path
    for path in sorted(Path(__file__).resolve().parent.glob("types_*.py")):
        importlib.import_module(path.stem)


_load_type_modules()
