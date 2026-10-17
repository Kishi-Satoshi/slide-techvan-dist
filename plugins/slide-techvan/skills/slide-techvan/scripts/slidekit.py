"""スライド生成の共通部品（図形・文字・はみ出し判定）。

build_pptx.py と build_template.py の両方から使う。
座標と大きさはすべてインチ（16:9 = 13.333 x 7.5 インチ）。
"""
import copy
import json
import re
import sys
import unicodedata
from pathlib import Path

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SKILL_DIR = Path(__file__).resolve().parent.parent
STYLES_DIR = SKILL_DIR / "styles"

SLIDE_W = 13.333
SLIDE_H = 7.5

FONT_JA = "游ゴシック"
FONT_EN = "Arial"
# 文字の高さに対する1行の高さ。游ゴシックは行の高さが大きいので余裕を見ておく
LINE_HEIGHT = 1.38

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


# ---------------------------------------------------------------- 実行環境

def console_utf8():
    """日本語版 Windows でも表示が文字化けしないよう、出力を UTF-8 にする。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass


def as_text(v):
    """None は空文字、数字などは文字列にする（deck.json の書き間違いで止まらないように）。"""
    if v is None:
        return ""
    return v if isinstance(v, str) else str(v)


# ---------------------------------------------------------------- 設定の読み込み

def load_theme(style):
    """styles/<style>/theme.json を読み込む。"""
    path = STYLES_DIR / style / "theme.json"
    if not path.exists():
        raise SystemExit(f"テーマが見つかりません: {path}")
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def rgb(hex_str):
    return RGBColor.from_string(hex_str.lstrip("#").upper())


# ---------------------------------------------------------------- 警告

class Warnings:
    """はみ出しや文字数超過の警告を集める。"""

    def __init__(self):
        self.items = []
        self.where = ""

    def add(self, msg):
        self.items.append(f"{self.where}{msg}")


WARN = Warnings()


# ---------------------------------------------------------------- 文字幅の見積もり

def char_em(ch, bold=False):
    """1文字の幅（em 単位）を見積もる。全角=1.0、半角英数≒0.6。"""
    if ch == "\t":
        return 2.0
    ea = unicodedata.east_asian_width(ch)
    if ea in ("F", "W", "A"):
        return 1.0
    if ch == " ":
        return 0.32
    if ch in "il.,:;|!'`":
        return 0.32
    if ch in "mwMW@%":
        return 0.9
    if ch.isupper() or ch.isdigit():
        return 0.66 if bold else 0.62
    return 0.58 if bold else 0.55


def text_em(text, bold=False, spacing_em=0.0):
    return sum(char_em(c, bold) + spacing_em for c in text)


def count_lines(text, width_in, size_pt, bold=False, spacing_em=0.0):
    """幅 width_in に size_pt で流し込んだときの行数を見積もる。"""
    if width_in <= 0:
        return 999
    width_em = width_in * 72.0 / size_pt
    lines = 0
    for para in text.split("\n"):
        if not para:
            lines += 1
            continue
        cur = 0.0
        n = 1
        for ch in para:
            w = char_em(ch, bold) + spacing_em
            if cur + w > width_em and cur > 0:
                # 行頭禁則（句読点・閉じ括弧は前の行に残す）は無視して単純に折り返す
                n += 1
                cur = w
            else:
                cur += w
        lines += n
    return lines


def text_height(text, width_in, size_pt, bold=False, line_spacing=1.0, spacing_em=0.0):
    n = count_lines(text, width_in, size_pt, bold, spacing_em)
    return n * size_pt * LINE_HEIGHT * line_spacing / 72.0


def fit_size(text, width_in, height_in, max_pt, min_pt, bold=False, line_spacing=1.0,
             spacing_em=0.0, label=""):
    """箱に収まる最大の文字サイズを返す。最小でも収まらなければ警告する。"""
    size = max_pt
    while size >= min_pt:
        if text_height(text, width_in, size, bold, line_spacing, spacing_em) <= height_in:
            return size
        size -= 0.5
    WARN.add(f"文字が多すぎて枠に収まりません（{label or text[:16]}…）。文を短くしてください")
    return min_pt


def max_chars_warn(text, limit, label):
    # 表示するときに取る空白（日本語と数字の間など）は数えない
    n = len(tidy_ja(as_text(text)).replace("\n", "").replace("**", ""))
    if n > limit:
        WARN.add(f"{label}が {n} 文字です（目安 {limit} 文字以内）")


# ---------------------------------------------------------------- 文字の装飾

EMPH_RE = re.compile(r"\*\*(.+?)\*\*")


def split_emphasis(text):
    """'**強調**' を [(文字, 強調か)] に分ける。"""
    parts = []
    pos = 0
    for m in EMPH_RE.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    return parts or [("", False)]


_CJK = r"\u3000-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uff00-\uffef"
_SP1 = re.compile(rf"(?<=[{_CJK}]) (?=[0-9A-Za-z%＋+\-−])")
_SP2 = re.compile(rf"(?<=[0-9A-Za-z%]) (?=[{_CJK}])")


_SP3 = re.compile(rf"(?<=[{_CJK}]) +(?=\*\*)|(?<=\*\*) +(?=[{_CJK}])")


def tidy_ja(text):
    """日本語と半角英数字の間の半角スペースを取る（「3,240 件」→「3,240件」）。
    強調（**…**）の前後で日本語と接するスペースも取る。"""
    return _SP2.sub("", _SP1.sub("", _SP3.sub("", text)))


def plain(text):
    return EMPH_RE.sub(r"\1", as_text(text))


def style_run(run, size=None, bold=None, color=None, font=None, latin=None, italic=None,
              spacing=None):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    ja = font or FONT_JA
    lat = latin or ja
    for tag, face in (("a:latin", lat), ("a:ea", ja), ("a:cs", lat)):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", face)
    # a:latin / a:ea / a:cs は solidFill より後ろに置く必要がある
    fill = rPr.find(qn("a:solidFill"))
    if fill is not None:
        rPr.remove(fill)
        rPr.insert(0, fill)
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    rPr.set("lang", "ja-JP")
    rPr.set("altLang", "en-US")


def set_paragraph(p, align="l", line_spacing=None, space_before=None, space_after=None):
    p.alignment = ALIGN[align]
    if line_spacing is not None:
        p.line_spacing = line_spacing
    if space_before is not None:
        p.space_before = Pt(space_before)
    if space_after is not None:
        p.space_after = Pt(space_after)


def write_runs(p, text, size, color, bold=False, emph_color=None, emph_bold=True, font=None,
               latin=None, spacing=None, italic=None, tidy=True):
    """段落に文字を書く。'**強調**' は emph_color（なければ太字のみ）で表示。"""
    text = as_text(text)
    if tidy:
        text = tidy_ja(text)
    for chunk, emph in split_emphasis(text):
        if not chunk:
            continue
        r = p.add_run()
        r.text = chunk
        style_run(
            r,
            size=size,
            bold=(emph_bold if emph else bold),
            color=(emph_color or color) if emph else color,
            font=font,
            latin=latin,
            spacing=spacing,
            italic=italic,
        )


# ---------------------------------------------------------------- 図形

def _no_shadow(shape):
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def add_box(shapes, x, y, w, h, fill=None, line=None, line_w=0.75, geom="rect", radius=None,
            shadow=False, dash=None, rotation=0, adjustments=None, name=None):
    """塗り・線つきの図形を置く。radius は角丸の半径（インチ）。"""
    shape_type = {
        "rect": MSO_SHAPE.RECTANGLE,
        "round": MSO_SHAPE.ROUNDED_RECTANGLE,
        "round_top": MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
        "oval": MSO_SHAPE.OVAL,
        "pill": MSO_SHAPE.ROUNDED_RECTANGLE,
        "chevron": MSO_SHAPE.CHEVRON,
        "pentagon": MSO_SHAPE.PENTAGON,
        "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
        "rtriangle": MSO_SHAPE.RIGHT_TRIANGLE,
        "arc": MSO_SHAPE.BLOCK_ARC,
        "wave": MSO_SHAPE.WAVE,
        "diamond": MSO_SHAPE.DIAMOND,
        "callout": MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
        "parallelogram": MSO_SHAPE.PARALLELOGRAM,
        "star": MSO_SHAPE.STAR_6_POINT,
        "plus": MSO_SHAPE.CROSS,
        "donut": MSO_SHAPE.DONUT,
        "tear": MSO_SHAPE.TEAR,
        "moon": MSO_SHAPE.MOON,
        "heart": MSO_SHAPE.HEART,
    }[geom]
    shp = shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    if geom in ("round", "round_top") and radius is not None:
        shp.adjustments[0] = max(0.0, min(0.5, radius / max(min(w, h), 0.01)))
    if geom == "pill":
        shp.adjustments[0] = 0.5
    if adjustments:
        for i, v in enumerate(adjustments):
            shp.adjustments[i] = v
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shp.line.dash_style = {"dash": MSO_LINE_DASH_STYLE.DASH,
                                   "dot": MSO_LINE_DASH_STYLE.ROUND_DOT}[dash]
    else:
        shp.line.fill.background()
    if rotation:
        shp.rotation = rotation
    if shadow:
        set_soft_shadow(shp)
    else:
        _no_shadow(shp)
    # 図形の中の文字の既定値
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return shp


def set_soft_shadow(shape, alpha=14, blur=0.06, dist=0.02, color="1F2937"):
    spPr = shape._element.spPr
    old = spPr.find(qn("a:effectLst"))
    if old is not None:
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(Emu(Inches(blur)))),
                          dist=str(int(Emu(Inches(dist)))), dir="5400000", algn="t",
                          rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val=color)
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def add_line(shapes, x1, y1, x2, y2, color, width=1.0, dash=None):
    ln = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = {"dash": MSO_LINE_DASH_STYLE.DASH,
                              "dot": MSO_LINE_DASH_STYLE.ROUND_DOT}[dash]
    return ln


def add_text(shapes, x, y, w, h, text, size=14, color="000000", bold=False, align="l",
             anchor="t", font=None, latin=None, line_spacing=1.0, spacing=None,
             emph_color=None, fit_min=None, label="", italic=None, para_space=0,
             margin=0.0, shape=None, wrap=True, bullet=None, bullet_color=None, tidy=True,
             marker=None):
    """文字を置く。text は文字列（改行で段落）か段落のリスト。

    marker に色を渡すと、1 行に収まるときに限り、**強調** の下に薄い帯（マーカー）を引く。

    fit_min を指定すると、枠に収まるよう size から fit_min まで縮める。
    shape を渡すとその図形の中に書く（新しい箱は作らない）。
    """
    paras = [as_text(t) for t in text] if isinstance(text, list) else as_text(text).split("\n")
    raw = "\n".join(plain(t) for t in paras)
    if shape is not None:
        # 既存の図形の中に書くときは、その図形の大きさで判定する
        w = shape.width / 914400
        h = shape.height / 914400
    inner_w = w - 2 * margin - (0.22 if bullet else 0)
    inner_h = h - 2 * margin
    if fit_min is not None:
        size = fit_size(raw, inner_w, inner_h, size, fit_min, bold, line_spacing,
                        (spacing / size) if spacing else 0.0, label)
    elif text_height(raw, inner_w, size, bold, line_spacing) > inner_h + 0.02 and label:
        WARN.add(f"「{label}」が枠からはみ出す可能性があります")
    if shape is None:
        shape = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = ANCHOR[anchor]
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for child in list(bodyPr):
        bodyPr.remove(child)
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_paragraph(p, align, line_spacing, space_before=para_space if i else 0)
        if bullet:
            _set_bullet(p, bullet, bullet_color or color, size)
        write_runs(p, t, size, color, bold=bold, emph_color=emph_color, font=font,
                   latin=latin, spacing=spacing, italic=italic, tidy=tidy)
    if marker and len(paras) == 1 and not bullet:
        sx = shape.left / 914400
        sy = shape.top / 914400
        add_marker_bars(shape, sx + margin, sy + margin, inner_w, inner_h, paras[0], size, bold,
                        align, anchor, marker, tidy=tidy)
    return shape


def add_marker_bars(text_shape, x, y, w, h, text, size, bold, align, anchor, color, tidy=True):
    """1 行の文の **強調** の下に、薄い色の帯（マーカー）を文字の後ろに引く。

    位置は文字幅の見積もり（text_em）で決める。2 行以上になる文には引かない。"""
    text = as_text(text)
    if tidy:
        text = tidy_ja(text)
    chunks = split_emphasis(text)
    if not any(e and c for c, e in chunks):
        return
    total = text_em(plain(text), bold) * size / 72.0
    if total > w * 0.98:          # 1 行に収まらない（折り返す）ときは引かない
        return
    line_h = size * LINE_HEIGHT / 72.0
    if align == "c":
        x0 = x + (w - total) / 2
    elif align == "r":
        x0 = x + w - total
    else:
        x0 = x
    if anchor == "m":
        top = y + (h - line_h) / 2
    elif anchor == "b":
        top = y + h - line_h
    else:
        top = y
    bar_top = top + line_h * MARKER_TOP
    bar_h = line_h * MARKER_H
    cx = x0
    for chunk, emph in chunks:
        cw = text_em(chunk, bold) * size / 72.0
        if emph and chunk.strip():
            pad = MARKER_PAD * size / 72.0     # 見積もりの誤差を吸収するため、左右に少し広げる
            bar = add_box(text_shape._parent, cx - pad, bar_top, cw + 2 * pad, bar_h, fill=color, geom="pill",
                          name="マーカー")
            text_shape._element.addprevious(bar._element)
        cx += cw


# マーカーの帯の位置（行の高さに対する割合）。文字の下の方に細く重なり、両端は丸い
MARKER_TOP = 0.68
MARKER_H = 0.18
MARKER_PAD = 0.15   # 左右に広げる幅（文字の大きさに対する割合）


def _set_bullet(p, char, color, size):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Emu(Inches(0.22)))))
    pPr.set("indent", str(int(-Emu(Inches(0.22)))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
        el = pPr.find(qn(tag))
        if el is not None:
            pPr.remove(el)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr"), val=color.lstrip("#").upper())
    etree.SubElement(pPr, qn("a:buSzPct"), val="100000")
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char=char)


def fill_placeholder(ph, text, size, color, bold=True, align="l", anchor=None, font=None,
                     latin=None, emph_color=None, fit_min=None, w=None, h=None, label="",
                     line_spacing=1.0, spacing=None, marker=None):
    """プレースホルダー（タイトル等）に文字を入れる。位置はレイアウトに従う。"""
    w = w if w is not None else ph.width / 914400
    h = h if h is not None else ph.height / 914400
    text = as_text(text)
    raw = plain(text)
    if fit_min is not None:
        size = fit_size(raw, w - 0.2, h - 0.1, size, fit_min, bold, line_spacing, 0.0,
                        label or raw[:16])
    tf = ph.text_frame
    tf.clear()
    if anchor:
        tf.vertical_anchor = ANCHOR[anchor]
    lines = text.split("\n")
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_paragraph(p, align, line_spacing)
        write_runs(p, t, size, color, bold=bold, emph_color=emph_color, font=font, latin=latin,
                   spacing=spacing)
    if marker and len(lines) == 1:
        lins = _inset(ph, "lIns", 91440)
        tins = _inset(ph, "tIns", 45720)
        anc = anchor or "m"
        add_marker_bars(ph, ph.left / 914400 + lins, ph.top / 914400 + tins, w - 2 * lins,
                        h - 2 * tins, text, size, bold, align, anc, marker)
    return size


def _inset(ph, attr, default):
    """プレースホルダーの内側の余白（インチ）。スライド → レイアウトの順に探す。"""
    shape = ph
    while shape is not None:
        bodyPr = shape._element.find(".//" + qn("a:bodyPr"))
        if bodyPr is not None and bodyPr.get(attr) is not None:
            return int(bodyPr.get(attr)) / 914400
        try:
            shape = shape._base_placeholder
        except AttributeError:
            shape = None
    return default / 914400


def spaced_ja(text):
    """「調査サマリー」→「調 査 サ マ リ ー」のように1文字ずつ空ける。"""
    return " ".join(list(plain(text)))


def place(shape, x=None, y=None, w=None, h=None):
    """プレースホルダーの位置を変える。継承している位置を明示してから変える
    （一部だけ指定すると、残りが 0 になってしまうため）。"""
    cur = (shape.left, shape.top, shape.width, shape.height)
    shape.left, shape.top, shape.width, shape.height = cur
    if x is not None:
        shape.left = Inches(x)
    if y is not None:
        shape.top = Inches(y)
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)


def clone_element(el):
    return copy.deepcopy(el)


# ---------------------------------------------------------------- 自由な形（モチーフ用）

def add_path(shapes, x, y, w, h, paths, fill=None, line=None, line_w=1.0, rotation=0,
             name=None, size=1000):
    """自由な形を置く。paths は 0〜size の座標で書いた
    [('M', x, y), ('L', x, y), ('C', x1, y1, x2, y2, x3, y3), ('Z',)] のリスト。
    複数の輪郭を描くときは paths をリストのリストにする。"""
    shp = add_box(shapes, x, y, w, h, fill=fill, line=line, line_w=line_w, rotation=rotation,
                  name=name)
    spPr = shp._element.spPr
    prst = spPr.find(qn("a:prstGeom"))
    cust = etree.Element(qn("a:custGeom"))
    for tag in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
        etree.SubElement(cust, qn(tag))
    etree.SubElement(cust, qn("a:rect"), l="0", t="0", r="r", b="b")
    pl = etree.SubElement(cust, qn("a:pathLst"))
    groups = paths if paths and isinstance(paths[0], list) else [paths]
    for cmds in groups:
        pe = etree.SubElement(pl, qn("a:path"), w=str(size), h=str(size))
        if fill is None:
            pe.set("fill", "none")
        for cmd in cmds:
            op = cmd[0]
            if op == "Z":
                etree.SubElement(pe, qn("a:close"))
                continue
            tag = {"M": "a:moveTo", "L": "a:lnTo", "C": "a:cubicBezTo"}[op]
            el = etree.SubElement(pe, qn(tag))
            pts = cmd[1:]
            for i in range(0, len(pts), 2):
                etree.SubElement(el, qn("a:pt"), x=str(int(pts[i])), y=str(int(pts[i + 1])))
    spPr.replace(prst, cust)
    return shp
